"""Read a workbook package in 1 pass: each sheet's cells with their saved values, formulas and number
formats, its merged cells, hidden rows and columns, comments, links and tables, and the workbook's defined
names and the other workbooks its formulas name. Reading formulas and saved values together takes 1 pass
where openpyxl takes 2, and empty cells are never stored.

Every part is checked on its way to the parser: a document type declaration is refused, and so is an
encoding that would not write one in ASCII and so could disguise it, so no entity is declared or expanded.
"""

import math
import posixpath
import re
from urllib.parse import unquote
from xml.etree.ElementTree import fromstring, iterparse

from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel, from_ISO8601

from .refs import MAX_COL, MAX_ROW, Area, parse_area

MAX_SPREAD = 100_000  # most cells a sheet's hyperlinks are spread over
_DOCTYPE = (b"<!DOCTYPE", "<!DOCTYPE".encode("utf-16-le"), "<!DOCTYPE".encode("utf-16-be"))
_ENCODING = re.compile(rb"""\s*<\?xml[^>]*?encoding\s*=\s*["']([^"']+)["']""")
_ASCII_SAFE = re.compile(rb"utf-?(8|16)(le|be)?|(us-)?ascii|iso[-_]?8859-\d+|latin-?\d|(windows|cp)-?125\d")
_LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"
_ESCAPED = re.compile(r"_x([0-9A-Fa-f]{4})_")
_THREADED = re.compile(r"^\[Threaded comment\].*?\n\s*Comment:\s*", re.S)
_TRUE = ("1", "true")


class PackageError(Exception):
    """A package this reader cannot take; the message follows the file's name."""


class SheetData:
    """One worksheet as stored. cells maps (row, col) to whatever make_cell built for it."""

    def __init__(self, name: str, state: str):
        self.name, self.state = name, state
        self.cells: dict = {}
        self.arrays: dict[tuple[int, int], Area] = {}  # the corner of an array formula or data table: its range
        self.merged: list[Area] = []
        self.hidden_rows: set[int] = set()
        self.hidden_cols: set[int] = set()
        self.comments: dict[tuple[int, int], str] = {}
        self.links: dict[tuple[int, int], str] = {}
        self.tables: list[tuple[str, Area, int, int, list[str]]] = []  # name, range, header and totals rows, columns


class Package:
    def __init__(self):
        self.sheets: list[SheetData] = []
        self.names: list[tuple[str, str, str | None, bool]] = []  # name, text, sheet it belongs to, hidden
        self.books: dict[str, str] = {}  # "1" in a formula's [1]Prices!A1: the file it names


def read_package(archive, make_cell, max_cells: int) -> Package:
    """make_cell(value, type code, formula or None, number format) builds what a sheet keeps for one cell."""
    return _Reader(archive, make_cell, max_cells).read()


class _Reader:
    def __init__(self, archive, make_cell, max_cells):
        self.archive, self.make_cell, self.max_cells = archive, make_cell, max_cells
        self.parts = {name.lower(): name for name in archive.namelist()}
        self.count = 0

    def read(self) -> Package:
        package = Package()
        office = [target for kind, target in self._rels("").values() if kind == "officeDocument"]
        book_path = office[0] if office else "xl/workbook.xml"
        root = self._xml(book_path)
        if root is None:
            raise PackageError("is not a workbook: it has no workbook part")
        rels = self._rels(book_path)
        epoch, entries, names, externals = CALENDAR_WINDOWS_1900, [], [], []
        for element in root.iter():
            tag = _local(element.tag)
            if tag == "workbookPr" and element.get("date1904") in _TRUE:
                epoch = CALENDAR_MAC_1904
            elif tag == "sheet":
                entries.append((element.get("name") or "", element.get("state") or "visible", _rid(element)))
            elif tag == "definedName":
                hidden = element.get("hidden") in _TRUE
                names.append((element.get("name") or "", element.text or "", element.get("localSheetId"), hidden))
            elif tag == "externalReference":
                externals.append(_rid(element))
        strings = self._strings(next((t for k, t in rels.values() if k == "sharedStrings"), "xl/sharedStrings.xml"))
        styles = self._styles(next((t for k, t in rels.values() if k == "styles"), "xl/styles.xml"))
        ns = root.tag[: root.tag.find("}") + 1]  # the spreadsheet namespace, transitional or strict, of every part
        for name, state, rid in entries:
            kind, path = rels.get(rid, ("", ""))
            if kind == "worksheet":
                package.sheets.append(self._sheet(SheetData(name, state), path, ns, strings, styles, epoch))
        order = [name for name, _, _ in entries]
        scoped = [
            (n, t, order[int(i)], h) for n, t, i, h in names if i is not None and i.isdigit() and int(i) < len(order)
        ]
        scoped.sort(key=lambda name: order.index(name[2]))
        package.names = scoped + [(n, t, None, h) for n, t, i, h in names if i is None]
        for number, rid in enumerate(externals, 1):
            kind, path = rels.get(rid, ("", ""))
            if kind == "externalLink":
                targets = [t for k, t in self._rels(path).values() if k in ("externalLinkPath", "xlPathMissing")]
                if targets:
                    package.books[str(number)] = posixpath.basename(unquote(targets[0].replace("\\", "/")))
        return package

    # Parts

    def _open(self, path):
        name = self.parts.get(path.lower())
        return _Checked(self.archive.open(name)) if name else None

    def _xml(self, path):
        name = self.parts.get(path.lower())
        if not name:
            return None
        data = self.archive.read(name)
        _check(data, first=True)
        return fromstring(data)

    def _rels(self, path) -> dict[str, tuple[str, str]]:
        """A part's relationships: id to (kind, the target's path in the package, or its link if external)."""
        folder, base = posixpath.split(path)
        root = self._xml(posixpath.join(folder, "_rels", base + ".rels"))
        found = {}
        for element in root.iter() if root is not None else ():
            if _local(element.tag) != "Relationship":
                continue
            target = (element.get("Target") or "").replace("\\", "/")
            if element.get("TargetMode") != "External":
                target = target[1:] if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
            found[element.get("Id")] = ((element.get("Type") or "").rsplit("/", 1)[-1], target)
        return found

    def _strings(self, path) -> list[str]:
        stream, strings = self._open(path), []
        if stream is None:
            return strings
        with stream:
            for _, element in iterparse(stream):
                if _local(element.tag) == "si":
                    strings.append(_text(element))
                    element.clear()
        return strings

    def _styles(self, path) -> list[tuple[str, bool, bool]]:
        """Each cell style's number format, and whether it shows a date or a duration."""
        root = self._xml(path)
        custom, ids = {}, []
        for part in root if root is not None else ():
            tag = _local(part.tag)
            if tag == "numFmts":
                custom = {int(f.get("numFmtId", -1)): f.get("formatCode") or "General" for f in part}
            elif tag == "cellXfs":
                ids = [int(xf.get("numFmtId") or 0) for xf in part]
        styles = []
        for number in ids:
            fmt = custom.get(number) or BUILTIN_FORMATS.get(number, "General")
            styles.append((fmt, is_date_format(fmt), is_timedelta_format(fmt)))
        return styles or [("General", False, False)]

    # A worksheet

    def _sheet(self, sheet, path, ns, strings, styles, epoch) -> SheetData:
        stream = self._open(path)
        if stream is None:
            raise PackageError(f"is damaged: sheet {sheet.name} is missing")
        ROW, COL, MERGE, LINK, TABLE = (ns + t for t in ("row", "col", "mergeCell", "hyperlink", "tablePart"))
        C, V, F, IS = (ns + t for t in ("c", "v", "f", "is"))
        shared, links, tables, hidden, merged, row = {}, [], [], [], {}, 0
        with stream:
            for _, element in iterparse(stream):
                tag = element.tag
                if tag == ROW:
                    number = element.get("r")
                    row = int(number) if number else row + 1
                    if element.get("hidden") in _TRUE and 0 < row <= MAX_ROW:
                        sheet.hidden_rows.add(row)
                    col = 0
                    for cell in element:
                        if cell.tag != C:
                            continue
                        ref = cell.get("r")
                        here, col = _coordinate(ref) if ref else (row, col + 1)
                        if not (0 < here <= MAX_ROW and 0 < col <= MAX_COL):
                            raise PackageError(f"is damaged: sheet {sheet.name} has a cell at {ref}, outside the grid")
                        value = formula = inline = None
                        for part in cell:
                            if part.tag == V:
                                value = part.text
                            elif part.tag == F:
                                formula = part
                            elif part.tag == IS:
                                inline = part
                        style = int(cell.get("s") or 0)
                        fmt, dated, elapsed = styles[style] if style < len(styles) else styles[0]
                        code = cell.get("t") or "n"
                        value, code = _value(value, code, inline, strings, dated, elapsed, epoch)
                        if formula is not None:
                            formula = self._formula(formula, here, col, shared, sheet)
                        elif value is None or isinstance(value, str) and not value.strip():
                            continue  # nothing shows in the cell
                        self.count += 1
                        if self.count > self.max_cells:
                            raise PackageError(f"has more than {self.max_cells:,} cells, more than this reader takes")
                        sheet.cells[(here, col)] = self.make_cell(value, code, formula, fmt)
                    element.clear()
                elif tag == COL and element.get("hidden") in _TRUE:
                    low = max(1, int(element.get("min") or 1))
                    hidden.append((low, min(MAX_COL, int(element.get("max") or low))))
                elif tag == MERGE and (area := parse_area(element.get("ref") or "")):
                    merged[area] = None
                elif tag == LINK:
                    links.append((element.get("ref") or "", _rid(element), element.get("location")))
                elif tag == TABLE:
                    tables.append(_rid(element))
        sheet.merged = list(merged)
        self._clear_merged(sheet)
        for low, high in _joined(hidden):
            sheet.hidden_cols.update(range(low, high + 1))
        rels, spread = self._rels(path), 0
        for ref, rid, location in links:
            target = rels.get(rid, ("", ""))[1] if rid else ""
            if not (area := parse_area(ref)) or not (target or location):
                continue
            spread += area.size
            if spread > MAX_SPREAD:  # past the budget, a link marks only its first cell
                area = Area(area.min_row, area.min_col, area.min_row, area.min_col)
            for r in range(area.min_row, area.max_row + 1):
                for c in range(area.min_col, area.max_col + 1):
                    sheet.links[(r, c)] = target or location
        for kind, target in rels.values():
            if kind == "comments":
                self._comments(sheet, target)
        for rid in tables:
            kind, target = rels.get(rid, ("", ""))
            if kind == "table":
                self._table(sheet, target)
        return sheet

    def _formula(self, element, row, col, shared, sheet):
        kind, text = element.get("t"), element.text
        formula = "=" + (text or "")
        if kind == "shared":
            key = element.get("si")
            if text:
                try:
                    shared[key] = Translator(formula, origin=_name(row, col))
                except Exception:  # the cells that share it then show no formula of their own
                    pass
                return formula
            try:
                return shared[key].translate_formula(_name(row, col)) if key in shared else formula
            except Exception:  # the translator rejects a formula it cannot tokenise
                return formula
        if kind in ("array", "dataTable") and (area := parse_area(element.get("ref") or "")):
            sheet.arrays[(row, col)] = area
            if kind == "dataTable":
                inputs = element.get("r1") or "", element.get("r2") or ""
                if element.get("dt2D") in _TRUE:
                    formula = f"=TABLE({inputs[0]},{inputs[1]})"
                else:
                    formula = f"=TABLE({inputs[0]},)" if element.get("dtr") in _TRUE else f"=TABLE(,{inputs[0]})"
        return formula

    def _clear_merged(self, sheet):
        """Keep only the first cell of each merged range, as Excel shows it, within a budget of work."""
        if not sheet.cells or not sheet.merged:
            return
        rows, cols = [r for r, _ in sheet.cells], [c for _, c in sheet.cells]
        used, work = Area(min(rows), min(cols), max(rows), max(cols)), 0
        for area in sheet.merged:
            corner, inside = (area.min_row, area.min_col), area.clip(used)
            if inside is None:
                continue
            work += min(inside.size, len(sheet.cells))
            if work > self.max_cells:
                break
            if inside.size <= len(sheet.cells):
                rows, cols = range(inside.min_row, inside.max_row + 1), range(inside.min_col, inside.max_col + 1)
                spots = [(r, c) for r in rows for c in cols]
            else:
                spots = [p for p in sheet.cells if inside.covers(*p)]
            for position in spots:
                if position != corner:
                    sheet.cells.pop(position, None)

    def _comments(self, sheet, path):
        root = self._xml(path)
        for element in root.iter() if root is not None else ():
            if _local(element.tag) == "comment" and (area := parse_area(element.get("ref") or "")):
                text = next((_text(part) for part in element if _local(part.tag) == "text"), "")
                sheet.comments[(area.min_row, area.min_col)] = _THREADED.sub("", text).strip()

    def _table(self, sheet, path):
        root = self._xml(path)
        if root is None or not (area := parse_area(root.get("ref") or "")):
            return
        columns = [e.get("name") or "" for e in root.iter() if _local(e.tag) == "tableColumn"]
        header, totals = int(root.get("headerRowCount") or 1), int(root.get("totalsRowCount") or 0)
        sheet.tables.append((root.get("displayName") or root.get("name") or "", area, header, totals, columns))


def _value(text, code, inline, strings, dated, elapsed, epoch):
    """A cell's saved value from its type code, and the code the reader keeps."""
    if code == "inlineStr":
        return (_text(inline) if inline is not None else None), "s"
    if not text:
        return None, code
    if code == "s":
        index = int(text)
        return (strings[index] if 0 <= index < len(strings) else None), "s"
    if code in ("str", "e"):
        return text, code
    if code == "b":
        return text.strip().lower() in _TRUE, "b"
    if code == "d":
        return from_ISO8601(text), "d"
    try:
        number = int(text) if text.lstrip("-").isdigit() else float(text)
    except ValueError:
        return text, "s"
    if not math.isfinite(number):
        return "#NUM!", "e"  # no spreadsheet saves infinity or not-a-number
    if dated:
        try:
            return from_excel(number, epoch, timedelta=elapsed), "d"
        except (OverflowError, ValueError):
            return "#VALUE!", "e"
    return number, "n"


def _text(element) -> str:
    """The text of a string item or comment: its plain text and runs, without phonetic guides."""
    parts = []
    for child in element:
        tag = _local(child.tag)
        if tag == "t":
            parts.append(child.text or "")
        elif tag == "r":
            parts += [part.text or "" for part in child if _local(part.tag) == "t"]
    return _ESCAPED.sub(_unescape, "".join(parts))


def _unescape(match):
    """_x000D_ is a carriage return; a code no text can hold, such as half of a surrogate pair, stays as written."""
    code = int(match[1], 16)
    return match[0] if 0xD800 <= code <= 0xDFFF else chr(code)


_COLUMNS: dict[str, int] = {}


def _coordinate(ref):
    """'BC12' as (12, 55). A reference no spreadsheet could write is damage."""
    split = len(ref) - len(ref.lstrip(_LETTERS))
    if not 0 < split <= 3 or len(ref) > 10 or not ref[split:].isdigit():
        raise PackageError(f"is damaged: a cell has the reference {ref[:12]}")
    letters = ref[:split].upper()
    if (col := _COLUMNS.get(letters)) is None:
        col = 0
        for letter in letters:
            col = col * 26 + ord(letter) - 64
        _COLUMNS[letters] = col
    return int(ref[split:]), col


def _joined(ranges):
    """Overlapping or touching ranges of numbers merged: (1, 3), (2, 5), (9, 9) as (1, 5), (9, 9)."""
    joined = []
    for low, high in sorted(ranges):
        if joined and low <= joined[-1][1] + 1:
            joined[-1][1] = max(joined[-1][1], high)
        else:
            joined.append([low, high])
    return joined


def _check(data, first):
    """Refuse XML that declares a document type, or that is in an encoding which could disguise one."""
    if any(marker in data for marker in _DOCTYPE):
        raise PackageError("has XML with a document type declaration, which is refused for safety")
    if first and (declared := _ENCODING.match(data)) and not _ASCII_SAFE.fullmatch(declared[1].lower()):
        encoding = declared[1].decode("ascii", "replace")[:20]
        raise PackageError(f"has XML in the {encoding} encoding, which is refused for safety")


class _Checked:
    """A part read as a stream, checked as it passes to the parser."""

    def __init__(self, stream):
        self.stream, self.tail = stream, None

    def read(self, size=-1):
        data = self.stream.read(size)
        _check((self.tail or b"") + data, first=self.tail is None)
        self.tail = ((self.tail or b"") + data)[-24:]  # a declaration split across 2 reads
        return data

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.stream.close()


def _name(row, col):
    letters = ""
    while col:
        col, rest = divmod(col - 1, 26)
        letters = chr(65 + rest) + letters
    return f"{letters}{row}"


def _local(tag) -> str:
    return tag.rpartition("}")[2]


def _rid(element):
    """An element's relationship id, in whichever namespace the package uses."""
    return next((value for key, value in element.attrib.items() if key.endswith("}id")), None)
