"""A workbook loaded once, then read in steps: outline, find, read, cell and trace."""

from datetime import date, datetime, time, timedelta
import difflib
from functools import cached_property
import hashlib
import io
from pathlib import Path
import re
from typing import NamedTuple
import zipfile
import zlib

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from . import formats
from .layout import Layout
from .reader import MAX_SPREAD, Package, PackageError, read_package
from .refs import Area, a1, cite, excel_formula, formula_refs, parse_area, quote, split_ref

SUFFIXES = (".xlsx", ".xlsm")
MAX_UNPACKED, MAX_CELLS = 512 * 2**20, 10_000_000
FIND_LIMIT, READ_CELLS, TRACE_DEPTH = 30, 600, 5
FOLLOWED = 12
_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_INTERNAL = ("_xlnm.", "_xlfn.", "_xlpm.")
_STRUCTURED = re.compile(r"([^\[\]]*)\[(.*)\]", re.S)
_ITEM = re.compile(r"\[((?:[^\[\]']|'.)*)\]")
_SCALES = (1e-3, 1e-6, 1e-9, 1e-12, 100)  # thousands, millions, billions and trillions, and percentages
_QUOTED_FIGURE = re.compile(r"(\d[\d,]*(?:\.\d+)?|\.\d+)")
_QUERY_NUMBER = re.compile(
    r"(?P<sign>[-+])?(?P<open>\()?[£$€¥]?(?P<whole>\d{1,3}(?:,\d{3})+|\d+)?(?P<fraction>\.\d+)?(?P<percent>%)?(?P<close>\))?"
)


class SheetError(Exception):
    """A file or reference that cannot be read. The message is written for the person or agent asking."""


class Cell:
    """A non-empty cell: its saved value, its formula and its number format."""

    __slots__ = ("value", "formula", "format", "kind", "saved", "_display")

    def __init__(self, value, formula, number_format, kind, saved):
        self.value, self.formula, self.format, self.kind, self.saved = value, formula, number_format, kind, saved
        self._display = None

    @property
    def display(self) -> str:
        """The text Excel shows in the cell."""
        if self._display is None:
            self._display = str(self.value) if self.kind == "error" else formats.display(self.value, self.format)
        return self._display


class Sheet:
    """One worksheet: its cells, merges, array formulas, hidden rows and columns, comments and links."""

    def __init__(self, name: str, state: str):
        self.name, self.state = name, state
        self.cells: dict[tuple[int, int], Cell] = {}
        self.arrays: dict[tuple[int, int], Area] = {}
        self.merged: list[Area] = []
        self.hidden_rows: set[int] = set()
        self.hidden_cols: set[int] = set()
        self.comments: dict[tuple[int, int], str] = {}
        self.links: dict[tuple[int, int], str] = {}

    @cached_property
    def layout(self) -> Layout:
        return Layout(self.cells, self.merged)

    @cached_property
    def bounds(self) -> Area | None:
        if not self.cells:
            return None
        rows, cols = [r for r, _ in self.cells], [c for _, c in self.cells]
        return Area(min(rows), min(cols), max(rows), max(cols))

    def hidden(self, row: int, col: int) -> str | None:
        if self.state != "visible":
            return "sheet"
        if row in self.hidden_rows:
            return "row"
        return "column" if col in self.hidden_cols else None

    def cells_in(self, area: Area | None) -> list[tuple[tuple[int, int], Cell]]:
        """The non-empty cells inside an area, row by row."""
        if area is None:
            return []
        if area.size <= len(self.cells):
            spots = (
                (r, c) for r in range(area.min_row, area.max_row + 1) for c in range(area.min_col, area.max_col + 1)
            )
            return [(p, self.cells[p]) for p in spots if p in self.cells]
        return sorted(((p, cell) for p, cell in self.cells.items() if area.covers(*p)), key=lambda item: item[0])


class Name(NamedTuple):
    name: str
    text: str
    scope: str | None
    hidden: bool


class Table(NamedTuple):
    name: str
    sheet: str
    area: Area
    header_rows: int
    totals_rows: int
    columns: list[str]


class Target(NamedTuple):
    """Cells a reference points to, and the defined name or table reference it went through."""

    sheet: str
    area: Area
    via: str | None = None


def load(path) -> "Book":
    """Read a workbook's formulas and saved values. Raises SheetError with a reason a reader can act on."""
    path = Path(path)
    if path.suffix.lower() not in SUFFIXES:
        raise SheetError(f"{path.name} is not an .xlsx or .xlsm workbook")
    try:
        raw = path.read_bytes()
    except FileNotFoundError:
        raise SheetError(f"No such workbook: {path.name}") from None
    except OSError as error:
        raise SheetError(f"Cannot read {path.name}: {error.strerror}") from None
    if path.name.startswith("~$") and not raw.startswith(b"PK"):
        raise SheetError(f"{path.name} is the lock file Excel keeps while {path.name[2:]} is open, not a workbook")
    if raw.startswith(_OLE):
        raise SheetError(
            f"{path.name} is password-protected or in the old .xls format; save it as .xlsx without a password"
        )
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
        _check_archive(archive, path.name)
    except SheetError:
        raise
    except Exception:  # zipfile and zlib raise several kinds of error on a damaged archive
        raise SheetError(f"{path.name} is damaged: it is not a readable .xlsx file") from None
    try:
        package = read_package(archive, _cell, MAX_CELLS)
    except PackageError as error:
        raise SheetError(f"{path.name} {error}") from None
    except (zipfile.BadZipFile, zlib.error, EOFError):
        raise SheetError(f"{path.name} is damaged: it is not a readable .xlsx file") from None
    except Exception as error:  # a damaged part or malformed XML
        raise SheetError(f"{path.name} could not be read: {type(error).__name__}: {error}") from None
    return Book(path, raw, package)


def _check_archive(archive, name):
    """Refuse an archive that unpacks too large. The reader checks each part's XML as it parses it."""
    unpacked = sum(item.file_size for item in archive.infolist())
    if unpacked > MAX_UNPACKED:
        raise SheetError(f"{name} unpacks to {unpacked / 2**20:,.0f} MB, over the {MAX_UNPACKED >> 20} MB limit")


def _cell(value, code, formula, number_format):
    saved = formula is None or value is not None or code == "str"
    return Cell(value, formula, number_format, _kind(value, code), saved)


class Book:
    """A loaded workbook. Each step returns a plain dict; render.markdown turns it into text."""

    def __init__(self, path, raw: bytes, package: Package):
        self.path = Path(path)
        self.name = self.path.name
        self.sha256 = hashlib.sha256(raw).hexdigest()
        self.sheets: dict[str, Sheet] = {}
        self.names: dict[tuple[str | None, str], Name] = {}
        self.tables: dict[str, Table] = {}
        self.books = package.books
        for data in package.sheets:
            self._add_sheet(data)
        for name, text, scope, hidden in package.names:
            if text and not name.startswith(_INTERNAL) and (scope is None or scope in self.sheets):
                self.names[(scope, name.casefold())] = Name(name, text, scope, hidden)
        self._by_name = {name.casefold(): sheet for name, sheet in self.sheets.items()}
        self._order = {name: i for i, name in enumerate(self.sheets)}
        self._row_free = not any("@" in n.text or "this row" in n.text.casefold() for n in self.names.values())
        self._inputs_of: dict = {}
        self._read_by: dict = {}  # what a formula's text reads on a sheet
        self._users_index: dict | None = None

    # Loading

    def _add_sheet(self, data):
        sheet = Sheet(data.name, data.state)
        sheet.cells, work = data.cells, 0
        for corner, area in data.arrays.items():
            inside = area.clip(sheet.bounds) if sheet.bounds else None
            work += min(inside.size, len(sheet.cells)) if inside else 0
            if inside is None or work > len(sheet.cells) + MAX_SPREAD:  # arrays that overlap are damage
                continue
            formula = sheet.cells[corner].formula if corner in sheet.cells else None
            for position, cell in sheet.cells_in(inside):
                cell.formula = cell.formula or formula  # cells an array formula fills show that formula
                sheet.arrays[position] = area
        sheet.merged, sheet.comments, sheet.links = data.merged, data.comments, data.links
        sheet.hidden_rows, sheet.hidden_cols = data.hidden_rows, data.hidden_cols
        for name, area, header, totals, columns in data.tables:
            self.tables[name.casefold()] = Table(name, sheet.name, area, header, totals, columns)
        self.sheets[sheet.name] = sheet

    # The reading steps

    def outline(self, max_blocks: int = 8, max_labels: int = 12, max_names: int = 50) -> dict:
        """Sheets with their blocks, header rows and row labels; defined names; tables; warnings.
        The more sheets and blocks a workbook has, the fewer blocks, labels and columns each one shows."""
        max_blocks = max(1, min(max_blocks, 64 // max(1, len(self.sheets))))
        shown = sum(min(len(sheet.layout.areas), max_blocks) for sheet in self.sheets.values())
        max_labels = max(3, min(max_labels, 240 // max(1, shown)))
        max_columns = max(4, min(20, 360 // max(1, shown)))
        sheets, unsaved, errors = [], 0, 0
        for sheet in self.sheets.values():
            cells = list(sheet.cells.values())
            info = {
                "name": sheet.name,
                "state": sheet.state,
                "range": sheet.bounds.a1() if sheet.bounds else "",
                "cells": len(cells),
                "formulas": sum(1 for cell in cells if cell.formula),
            }
            missing, wrong = sum(1 for c in cells if not c.saved), sum(1 for c in cells if c.kind == "error")
            unsaved, errors = unsaved + missing, errors + wrong
            extras = {
                "unsaved": missing,
                "errors": wrong,
                "comments": len(sheet.comments),
                "hidden_rows": _spans(sheet.hidden_rows, str),
                "hidden_columns": _spans(sheet.hidden_cols, get_column_letter),
            }
            info.update({key: value for key, value in extras.items() if value})
            count = len(sheet.layout.areas)
            largest = sorted(range(count), key=lambda b: -sheet.layout.size(b))[:max_blocks]
            info["blocks"] = [self._block(sheet, b, max_labels, max_columns) for b in sorted(largest)]
            if count > max_blocks:
                info["more_blocks"] = count - max_blocks
            sheets.append(info)
        notes = []
        if unsaved:
            subject = "1 formula has" if unsaved == 1 else f"{unsaved:,} formulas have"
            notes.append(
                f"{subject} no saved value, so they read as empty: the file was last saved by a program that "
                "does not calculate. Open and save it in Excel to store the values."
            )
        if errors:
            subject = "1 cell shows" if errors == 1 else f"{errors:,} cells show"
            notes.append(f"{subject} an error such as #REF! or #DIV/0!.")
        first = next(iter(self.sheets), None)
        names = [name for name in self.names.values() if not name.hidden]
        hidden = len(self.names) - len(names)
        result = {
            "kind": "outline",
            "workbook": self.name,
            "sha256": self.sha256,
            "cite_as": self.citation(first, "A1") if first else "",
            "sheets": sheets,
            "names": [self._name_info(name) for name in names[:max_names]],
            "tables": [
                {"name": t.name, "sheet": t.sheet, "range": t.area.a1(), "columns": t.columns}
                for t in self.tables.values()
            ],
            "warnings": notes,
        }
        if len(names) > max_names:
            result["more_names"] = len(names) - max_names
        if hidden:
            result["hidden_names"] = hidden
        return result

    def find(self, query: str, sheet: str | None = None, formulas: bool = False, limit: int = FIND_LIMIT) -> dict:
        """Cells whose text or shown value contains the query, or whose number, as stored or as shown, rounds
        to it: "12.5%" finds 0.1249, and "1,597" finds 1,597,000 shown in thousands."""
        needle = query.strip().casefold()
        if not needle:
            raise SheetError("Give some text or a number to find")
        number = None if formulas else _query_number(needle)
        sheets = [self._sheet_or_error(sheet)] if sheet else list(self.sheets.values())
        hits = []
        for order, page in enumerate(sheets):
            for (row, col), cell in page.cells.items():
                rank = _rank(cell, needle, number, formulas)
                if rank is not None:
                    hits.append((rank, order, row, col))
        hits.sort()
        shown = hits[: max(1, min(limit, 500))]
        return {
            "kind": "find",
            "workbook": self.name,
            "sha256": self.sha256,
            "query": query.strip(),
            "searched": "formulas" if formulas else "values",
            "total": len(hits),
            "matches": [self._info(sheets[order], row, col) for _, order, row, col in shown],
        }

    def read(self, ref: str, max_cells: int = READ_CELLS) -> dict:
        """A range as rows and columns, with the labels and headers just outside it and its formulas."""
        sheet, area = self.locate(ref)
        shown = area.clip(sheet.bounds) if sheet.bounds else None
        span = shown or area
        found = sheet.cells_in(shown)
        total, budget = len(found), max(1, min(max_cells, 20_000))
        if cols := sorted({p[1] for p, _ in found})[:budget]:
            found = [(p, cell) for p, cell in found if p[1] <= cols[-1]]
            rows = sorted({p[0] for p, _ in found})[: max(1, budget // len(cols))]
            found = [(p, cell) for p, cell in found if p[0] <= rows[-1]]
        rows, cols = sorted({p[0] for p, _ in found}), sorted({p[1] for p, _ in found})
        truncated = len(found) < total
        layout = sheet.layout
        first_in_column, first_in_row = {}, {}
        for position, _ in found:
            first_in_column.setdefault(position[1], position)
            first_in_row.setdefault(position[0], position)
        columns = []
        for col in cols:
            item = {"col": get_column_letter(col)}
            heads = layout.header(*first_in_column[col])
            if heads and heads[-1][0] < span.min_row:
                item["header"] = _joined(sheet, heads)
            if col in sheet.hidden_cols:
                item["hidden"] = True
            columns.append(item)
        row_items = []
        for row in rows:
            item = {"row": row}
            if labels := layout.label(row, span.min_col, layout.block[first_in_row[row]]):
                item["label"] = _joined(sheet, labels)
            if row in sheet.hidden_rows:
                item["hidden"] = True
            row_items.append(item)
        result = {
            "kind": "read",
            "workbook": self.name,
            "sha256": self.sha256,
            "sheet": sheet.name,
            "range": span.a1(),
            "citation": self.citation(sheet.name, span.a1()),
            "columns": columns,
            "rows": row_items,
            "cells": [self._info(sheet, row, col, context=False) for (row, col), _ in found],
            "formulas": _formula_groups({p: cell.formula for p, cell in found if cell.formula}, sheet.arrays),
            "merged": [m.a1() for m in sheet.merged if shown and m.clip(shown)],
            "truncated": truncated,
        }
        if sheet.state != "visible":
            result["hidden"] = "sheet"
        return result

    def cell(self, ref: str) -> dict:
        """One cell in full: value, format, formula, labels, the cells it reads and the cells that read it."""
        sheet, row, col = self._one_cell(ref)
        info = self._info(sheet, row, col)
        cell = sheet.cells.get((row, col))
        if cell is not None:
            info["number_format"] = cell.format
        if merged := next((m for m in sheet.merged if m.covers(row, col)), None):
            info["merged"] = merged.a1()
        if comment := sheet.comments.get((row, col)):
            info["comment"] = comment
        if link := sheet.links.get((row, col)):
            info["link"] = link
        if cell is not None and cell.formula:
            targets, notes = self._inputs(sheet.name, (row, col))
            info["precedents"] = [self._target(target) for target in targets]
            if notes:
                info["notes"] = notes
        users = self._users(sheet.name, (row, col))
        info["dependents"] = [_short(name, a1(*position)) for name, position in users[:50]]
        if len(users) > 50:
            info["dependents_total"] = len(users)
        return {"kind": "cell", "workbook": self.name, "sha256": self.sha256, **info}

    def trace(self, ref: str, dependents: bool = False, depth: int = TRACE_DEPTH, max_nodes: int = 200) -> dict:
        """The tree of cells behind one cell's formula, down to its inputs; or the formulas that use it."""
        sheet, row, col = self._one_cell(ref)
        depth, budget, seen = max(1, min(depth, 20)), [max_nodes], set()

        def visit(name, position, via, level, path):
            budget[0] -= 1
            node = self._node(name, position, via)
            cell = self.sheets[name].cells.get(position)
            key = (name, position)
            users = self._users(name, position) if dependents else []
            if not (users if dependents else cell is not None and cell.formula):
                return node
            if key in path:
                return {**node, "circular": True}
            if key in seen:
                return {**node, "repeat": True}
            seen.add(key)
            if level >= depth or budget[0] <= 0:
                return {**node, "more": True}
            if dependents:
                steps, notes = [Target(user, Area(*place, *place)) for user, place in users], []
            else:
                steps, notes = self._inputs(name, position)
            children = []
            for step in steps:
                if budget[0] <= 0:
                    node["more"] = True
                    break
                if step.area.size == 1:
                    place = (step.area.min_row, step.area.min_col)
                    children.append(visit(step.sheet, place, step.via, level + 1, path | {key}))
                    continue
                budget[0] -= 1
                page = self.sheets[step.sheet]
                inside = page.cells_in(step.area.clip(page.bounds) if page.bounds else None)
                group = self._range_node(step, inside)
                members = [place for place, member in inside if member.formula]
                followed = [visit(step.sheet, place, None, level + 1, path | {key}) for place in members[:FOLLOWED]]
                if followed:
                    group["children"] = followed
                if len(members) > FOLLOWED:
                    group["skipped"] = len(members) - FOLLOWED
                children.append(group)
            children += [{"note": note} for note in notes]
            if children:
                node["children"] = children
            return node

        root = visit(sheet.name, (row, col), None, 0, frozenset())
        direction = "dependents" if dependents else "precedents"
        return {"kind": "trace", "workbook": self.name, "sha256": self.sha256, "direction": direction, "root": root}

    # References

    def citation(self, sheet: str, ref: str) -> str:
        return cite(self.name, sheet, ref)

    def locate(self, ref: str) -> tuple[Sheet, Area]:
        """The sheet and area a reference names: a range, a sheet, a defined name, a table or a citation."""
        text = ref.strip()
        if not text:
            raise SheetError("Give a reference such as Model!B4, a sheet name, a defined name or a table")
        if not self.sheets:
            raise SheetError(f"{self.name} has no worksheets")
        book, sheet_name, rest = split_ref(text)
        if book is not None:
            if book.casefold() != self.name.casefold():
                raise SheetError(f"{ref} cites {book}, but this workbook is {self.name}")
            text = f"{quote(sheet_name)}!{rest}" if sheet_name else rest
        bare = text[1:-1].replace("''", "'") if len(text) > 1 and text[0] == text[-1] == "'" else text
        whole = self._sheet(sheet_name) if sheet_name is not None and not rest else self._sheet(bare)
        if whole is not None:
            return whole, whole.bounds or Area(1, 1, 1, 1)
        default = next(iter(self.sheets)) if len(self.sheets) == 1 else None
        if sheet_name is None and parse_area(rest) and default is None:
            example = f"{quote(next(iter(self.sheets)))}!{rest.upper()}"
            raise SheetError(f"Name the sheet, for example {example}. Sheets: {', '.join(self.sheets)}")
        targets, notes = self._resolve(text, default)
        if len(targets) == 1:
            return self.sheets[targets[0].sheet], targets[0].area
        if targets:
            areas = ", ".join(_short(t.sheet, t.area.a1()) for t in targets)
            raise SheetError(f"{ref} covers {len(targets)} areas ({areas}); read them one at a time")
        raise SheetError(self._unknown(ref, text, notes))

    def _one_cell(self, ref):
        sheet, area = self.locate(ref)
        if area.size != 1:
            raise SheetError(f"{ref} is the range {area.a1()}; give one cell, or use read for a range")
        return sheet, area.min_row, area.min_col

    def _resolve(self, text, sheet, at=None, via=None, depth=0):
        """The cells one reference points to, and notes on anything that cannot be followed."""
        book, sheet_part, ref = split_ref(text)
        if book is not None:
            return [], [f"{text} is in another workbook ({self.books.get(book, book)}), so it is not followed"]
        if sheet_part is not None:
            first, _, last = sheet_part.partition(":")
            span = self._span(first, last or first)
            if span is None:
                return [], [f"{text}: there is no sheet called {sheet_part}"]
            if area := parse_area(ref):
                return [Target(page.name, area, via) for page in span], []
            if len(span) == 1 and (name := self.names.get((span[0].name, ref.casefold()))):
                return self._name_targets(name, span[0].name, at, via, depth)
            return [], [f"{text} is not a cell reference or a name on {sheet_part}"]
        if area := parse_area(ref):
            return ([Target(sheet, area, via)], []) if sheet else ([], [f"{text} needs a sheet name"])
        if "[" in ref:
            return self._structured(ref, sheet, at, via)
        key = ref.casefold()
        if name := self.names.get((sheet, key)) or self.names.get((None, key)):
            return self._name_targets(name, sheet, at, via, depth)
        if table := self.tables.get(key):
            return self._structured(f"{table.name}[]", sheet, at, via)
        return [], [f"{text} is not a sheet, defined name or table in {self.name}"]

    def _name_targets(self, name, sheet, at, via, depth):
        if depth > 8:
            return [], [f"{name.name} refers to itself or nests too deeply"]
        if "#REF!" in name.text.upper():
            return [], [f"{name.name} refers to a deleted range: ={name.text}"]
        refs, warning = formula_refs("=" + name.text)
        targets, notes = [], [f"{name.name} {warning}"] if warning else []
        if not refs:
            notes.append(f"{name.name} is the constant ={name.text}")
        for ref in refs:
            found, problems = self._resolve(ref, name.scope or sheet, at, via or name.name, depth + 1)
            targets += found
            notes += problems
        return targets, notes

    def _structured(self, ref, sheet, at, via):
        """A table reference such as Sales[Amount], Sales[#All] or Sales[[#This Row],[Amount]]."""
        match = _STRUCTURED.fullmatch(ref)
        if not match:
            return [], [f"{ref} is not a table reference this reader can follow"]
        name, spec = match[1], match[2].strip()
        table = self.tables.get(name.casefold()) if name else self._table_at(sheet, at)
        if table is None:
            return [], [f"{ref}: there is no table called {name}" if name else f"{ref}: this cell is not in a table"]
        this_row = spec.startswith("@")
        spec = spec[1:].strip() if this_row else spec
        items = [item.strip() for item in _ITEM.findall(spec)] if spec.startswith("[") else [spec] if spec else []
        specials = {item.casefold() for item in items if item.startswith("#")}
        columns = [_unescape(item) for item in items if not item.startswith("#")]
        area = table.area
        first, last = area.min_row + table.header_rows, area.max_row - table.totals_rows
        if this_row or "#this row" in specials:
            if at is None or not first <= at[0] <= last:
                return [], [f"{ref} means this row of {table.name}, but the cell is not in its rows"]
            rows = (at[0], at[0])
        else:
            spans = {
                "#all": (area.min_row, area.max_row),
                "#data": (first, last),
                "#headers": (area.min_row, first - 1),
                "#totals": (last + 1, area.max_row),
            }
            chosen = [spans[s] for s in specials if s in spans] or [spans["#data"]]
            rows = (min(r for r, _ in chosen), max(r for _, r in chosen))
            if rows[0] > rows[1]:
                return [], [f"{ref}: {table.name} has no such rows"]
        known = [column.casefold() for column in table.columns]
        if missing := [c for c in columns if c.casefold() not in known]:
            return [], [f"{ref}: {table.name} has no column called {missing[0]}"]
        picked = [known.index(c.casefold()) for c in columns] or [0, len(known) - 1]
        cols = (area.min_col + min(picked), area.min_col + max(picked))
        return [Target(table.sheet, Area(rows[0], cols[0], rows[1], cols[1]), via or ref)], []

    def _inputs(self, sheet: str, position) -> tuple[list[Target], list[str]]:
        """What a formula cell reads, resolved through names and tables, plus notes on what was not followed.
        A formula's text names the same cells wherever it sits, unless it reads its own row of a table."""
        key = (sheet, position)
        if key not in self._inputs_of:
            formula = self.sheets[sheet].cells[position].formula
            fixed = self._row_free and "@" not in formula and "this row" not in formula.casefold()
            if not fixed or (sheet, formula) not in self._read_by:
                refs, warning = formula_refs(formula)
                targets, notes = [], [f"The formula {warning}."] if warning else []
                for ref in refs:
                    found, problems = self._resolve(ref, sheet, position)
                    targets += [target for target in found if target not in targets]
                    notes += [note for note in problems if note not in notes]
                if not fixed:
                    self._inputs_of[key] = (targets, notes)
                    return self._inputs_of[key]
                self._read_by[(sheet, formula)] = (targets, notes)
            self._inputs_of[key] = self._read_by[(sheet, formula)]
        return self._inputs_of[key]

    def _users(self, sheet: str, position) -> list[tuple[str, tuple[int, int]]]:
        """The formula cells that read a cell, in workbook order. The index is built on first use: what reads
        each single cell, then ranges by column, then ranges too wide to file by column."""
        if self._users_index is None:
            index = {}
            for page in self.sheets.values():
                for place, cell in page.cells.items():
                    if cell.formula:
                        for target in self._inputs(page.name, place)[0]:
                            cells, columns, wide = index.setdefault(target.sheet, ({}, {}, []))
                            user, area = (page.name, place), target.area
                            if area.size == 1:
                                cells.setdefault((area.min_row, area.min_col), []).append(user)
                            elif area.max_col - area.min_col < 64:
                                for col in range(area.min_col, area.max_col + 1):
                                    columns.setdefault(col, []).append((area, user))
                            else:
                                wide.append((area, user))
            self._users_index = index
        if sheet not in self._users_index:
            return []
        cells, columns, wide = self._users_index[sheet]
        found = set(cells.get(position, ()))
        found.update(user for area, user in columns.get(position[1], ()) if area.covers(*position))
        found.update(user for area, user in wide if area.covers(*position))
        return sorted(found, key=lambda user: (self._order[user[0]], user[1]))

    # Helpers

    def _sheet(self, name):
        return self._by_name.get(name.casefold())

    def _sheet_or_error(self, name):
        if sheet := self._sheet(name):
            return sheet
        raise SheetError(self._unknown(name, name, [f"There is no sheet called {name}"]))

    def _span(self, first, last):
        start, end = self._sheet(first), self._sheet(last)
        if start is None or end is None:
            return None
        if start is end:
            return [start]
        i, j = sorted((self._order[start.name], self._order[end.name]))
        return list(self.sheets.values())[i : j + 1]

    def _table_at(self, sheet, at):
        return next((t for t in self.tables.values() if t.sheet == sheet and at and t.area.covers(*at)), None)

    def _unknown(self, ref, text, notes):
        _, sheet_part, rest = split_ref(text)
        word = (sheet_part if sheet_part and not self._sheet(sheet_part) else rest).casefold()
        pool = {name.casefold(): name for name in self.sheets}
        pool |= {key: name.name for (scope, key), name in self.names.items() if scope is None}
        pool |= {key: table.name for key, table in self.tables.items()}
        close = [] if word in pool else [pool[m] for m in difflib.get_close_matches(word, pool, n=3, cutoff=0.6)]
        reason = notes[0] if notes else f"Nothing called {ref} in {self.name}"
        return reason + (f". Did you mean {' or '.join(close)}?" if close else "")

    def _info(self, sheet, row, col, context=True) -> dict:
        ref = a1(row, col)
        info = {"sheet": sheet.name, "cell": ref, "citation": self.citation(sheet.name, ref)}
        cell = sheet.cells.get((row, col))
        if cell is None:
            info.update(value=None, display="", type="empty")
        else:
            info.update(value=_plain(cell.value), display=cell.display, type=cell.kind)
            if cell.formula:
                info["formula"] = excel_formula(cell.formula)
                if not cell.saved:
                    info["saved"] = False
        if array := sheet.arrays.get((row, col)):
            info["array"] = array.a1()
        if context:
            info.update(self._context(sheet, row, col))
        if hidden := sheet.hidden(row, col):
            info["hidden"] = hidden
        return info

    def _context(self, sheet, row, col) -> dict:
        found = {"label": sheet.layout.label(row, col), "header": sheet.layout.header(row, col)}
        return {key: _joined(sheet, positions) for key, positions in found.items() if positions}

    def _node(self, sheet_name, position, via=None) -> dict:
        sheet = self.sheets[sheet_name]
        cell = sheet.cells.get(position)
        node = {"ref": _short(sheet_name, a1(*position)), "display": cell.display if cell else ""}
        if via:
            node["via"] = via
        node.update(self._context(sheet, *position))
        if cell is not None and cell.formula:
            node["formula"] = excel_formula(cell.formula)
            if not cell.saved:
                node["saved"] = False
        return node

    def _target(self, target: Target) -> dict:
        if target.area.size == 1:
            return self._node(target.sheet, (target.area.min_row, target.area.min_col), target.via)
        sheet = self.sheets[target.sheet]
        return self._range_node(target, sheet.cells_in(target.area.clip(sheet.bounds) if sheet.bounds else None))

    def _range_node(self, target: Target, inside) -> dict:
        node = {"ref": _short(target.sheet, target.area.a1())}
        if target.via:
            node["via"] = target.via
        node["cells"] = len(inside)
        node["formulas"] = sum(1 for _, cell in inside if cell.formula)
        if len(inside) <= 12:
            node["values"] = [cell.display for _, cell in inside]
        return node

    def _block(self, sheet, block, max_labels, max_columns):
        layout = sheet.layout
        area = layout.areas[block]
        rows = layout.rows_of(block)
        info = {"range": area.a1()}
        start = rows[0]
        if heads := layout.header_rows(block):
            top = heads[0]
            if rows[0] < layout.band(block, top)[0]:
                first = next(c for c in layout.rows[rows[0]] if layout.block[(rows[0], c)] == block)
                info["title"] = sheet.cells[(rows[0], first)].display
            columns = []
            for col in range(area.min_col, area.max_col + 1):
                if stack := layout.headings(block, top, col):
                    columns.append({"col": get_column_letter(col), "text": _joined(sheet, stack)})
            info["header_row"] = top
            info["columns"] = columns[:max_columns]
            if len(columns) > max_columns:
                info["more_columns"] = len(columns) - max_columns
            start = top + 1
        labels = []
        for row in rows:
            if row >= start and (positions := layout.label(row, area.max_col + 1, block)):
                labels.append({"row": row, "text": _joined(sheet, positions)})
        info["labels"] = labels[:max_labels]
        if len(labels) > max_labels:
            info["more_labels"] = len(labels) - max_labels
        return info

    def _name_info(self, name: Name) -> dict:
        info = {"name": name.name, "refers_to": name.text}
        if name.scope:
            info["scope"] = name.scope
        if name.hidden:
            info["hidden"] = True
        targets, _ = self._name_targets(name, name.scope, None, None, 0)
        if len(targets) == 1 and targets[0].area.size == 1:
            target = targets[0]
            if cell := self.sheets[target.sheet].cells.get((target.area.min_row, target.area.min_col)):
                info["value"] = cell.display
        return info


def _kind(value, code):
    if value is None or isinstance(value, str) and not value.strip():
        return "empty"
    if code == "e":
        return "error"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime, date, time, timedelta)):
        return "date"
    return "text"


def _plain(value):
    """A value JSON can hold: dates and times become ISO 8601 text."""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    return value


def _joined(sheet, positions):
    return " / ".join(sheet.cells[position].display for position in positions)


def _short(sheet, ref):
    return f"{quote(sheet)}!{ref}"


def _unescape(text):
    return re.sub(r"'(.)", r"\1", text).strip()


def _spans(numbers, name):
    """Group numbers into runs, e.g. {3, 4, 5, 9} as "3:5, 9"."""
    runs = []
    for number in sorted(numbers):
        if runs and number == runs[-1][1] + 1:
            runs[-1][1] = number
        else:
            runs.append([number, number])
    return ", ".join(name(a) if a == b else f"{name(a)}:{name(b)}" for a, b in runs)


def _query_number(text):
    """A number in a query, with the rounding it allows: "12.5%" is 0.125 give or take 0.0005."""
    match = _QUERY_NUMBER.fullmatch(text.replace(" ", ""))
    if not match or not (match["whole"] or match["fraction"]) or bool(match["open"]) != bool(match["close"]):
        return None
    fraction = match["fraction"] or ""
    value = float((match["whole"] or "0").replace(",", "") + fraction)
    tolerance = 0.5 * 10 ** -max(len(fraction) - 1, 0)
    if match["sign"] == "-" or match["open"]:
        value = -value
    if match["percent"]:
        value, tolerance = value / 100, tolerance / 100
    return value, tolerance * (1 + 1e-9)


def _rank(cell, needle, number, formulas):
    """0 for an exact match, 1 for a number that rounds to the query, 2 for a partial match; None for no match."""
    if formulas:
        return 2 if cell.formula and needle in excel_formula(cell.formula).casefold() else None
    if cell.kind == "number":
        if number is not None:
            return _rank_number(cell, needle, number)
        if not formats.may_show(needle, cell.format):
            return None
    shown = cell.display.casefold()
    text = cell.value.casefold() if isinstance(cell.value, str) else ""
    if needle in (shown, text):
        return 0
    return 2 if needle in shown or (text and needle in text) else None


def _rank_number(cell, needle, number):
    """0 when a number cell shows the query, 1 when its value or the number it shows rounds to it: 1,597,000
    shown in thousands as 1,597 matches "1,597". Digits inside a longer number do not match."""
    target, tolerance = number
    if abs(cell.value - target) <= tolerance:
        return 0 if cell.display.casefold() == needle else 1
    if not any(abs(cell.value * scale - target) <= tolerance + 0.5 for scale in _SCALES):
        return None  # no format could show a number this close to the query
    shown = cell.display.casefold()
    if shown == needle:
        return 0
    figure = _QUOTED_FIGURE.search(shown)
    if figure is None:
        return None
    value = float(figure[1].replace(",", ""))
    value = -value if "(" in shown[: figure.start()] or "-" in shown[: figure.start()] else value
    return 1 if abs(value - target) <= tolerance else None


def _formula_groups(formulas: dict, arrays: dict) -> list[dict]:
    """Formulas grouped by the rectangles they were filled across, each shown once for its first cell."""
    done, groups = set(), []
    for start in sorted(formulas):
        if start in done:
            continue
        formula = formulas[start]
        if start in arrays:
            members = [p for p in formulas if arrays.get(p) == arrays[start]]
            done.update(members)
            rows, cols = [r for r, _ in members], [c for _, c in members]
            area = Area(min(rows), min(cols), max(rows), max(cols))
            groups.append({"range": area.a1(), "formula": excel_formula(formula), "array": True})
            continue
        try:
            translator = Translator(formula, origin=a1(*start))
        except Exception:  # the translator's tokenizer rejects malformed formulas
            translator = None

        def same(position):
            if translator is None or position in done or position in arrays or position not in formulas:
                return False
            try:
                return translator.translate_formula(a1(*position)) == formulas[position]
            except Exception:
                return False

        row, col = start
        right = col
        while same((row, right + 1)):
            right += 1
        bottom = row
        while all(same((bottom + 1, c)) for c in range(col, right + 1)):
            bottom += 1
        done.update((r, c) for r in range(row, bottom + 1) for c in range(col, right + 1))
        groups.append({"range": Area(row, col, bottom, right).a1(), "formula": excel_formula(formula)})
    return groups
