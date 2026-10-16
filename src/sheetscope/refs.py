"""References: A1 areas, quoted sheet names, citations and the references a formula reads."""

from functools import lru_cache
import re
from typing import NamedTuple

from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

MAX_ROW, MAX_COL = 1_048_576, 16_384

_CELL = r"\$?([A-Za-z]{1,3})\$?([0-9]+)"
_AREA = re.compile(rf"{_CELL}(?::{_CELL})?")
_COLUMNS = re.compile(r"\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})")
_ROWS = re.compile(r"\$?([0-9]+):\$?([0-9]+)")
_PLAIN = re.compile(r"[A-Za-z_][A-Za-z0-9_.]*")
_CELL_LIKE = re.compile(r"[A-Za-z]{1,3}[0-9]+|[RrCc][0-9]*|[Rr][0-9]*[Cc][0-9]*")
_BOOK = re.compile(r"[A-Za-z0-9_.]+")
_SPILL = re.compile(r"(\$?[A-Za-z]{1,3}\$?[0-9]+)#")
_STRING = re.compile(r'"(?:[^"]|"")*"')
_QUOTED = re.compile(r'("(?:[^"]|"")*")')
_DYNAMIC = {"INDIRECT(", "OFFSET("}
_ONE = r"(?:'(?:[^']|'')+'!|[A-Za-z0-9_.]+!)?\$?[A-Za-z]{1,3}\$?[0-9]+"
_ANCHOR = re.compile(rf"_xlfn\.ANCHORARRAY\(({_ONE})\)", re.I)
_SINGLE = re.compile(rf"_xlfn\.SINGLE\(({_ONE}(?::\$?[A-Za-z]{{1,3}}\$?[0-9]+)?)\)", re.I)
_PREFIX = re.compile(r"_xl(?:fn|ws|pm|udf)\.", re.I)
_CELLISH = re.compile(r"(?<![A-Za-z0-9_.$\\])\$?[A-Za-z]{1,3}\$?[0-9]+(?![A-Za-z0-9_.(])")


class Area(NamedTuple):
    """A rectangle of cells, 1-based and inclusive."""

    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def size(self) -> int:
        return (self.max_row - self.min_row + 1) * (self.max_col - self.min_col + 1)

    def covers(self, row: int, col: int) -> bool:
        return self.min_row <= row <= self.max_row and self.min_col <= col <= self.max_col

    def clip(self, other: "Area") -> "Area | None":
        """The part of this area inside another, or None when they do not overlap."""
        top, left = max(self.min_row, other.min_row), max(self.min_col, other.min_col)
        bottom, right = min(self.max_row, other.max_row), min(self.max_col, other.max_col)
        return Area(top, left, bottom, right) if top <= bottom and left <= right else None

    def a1(self) -> str:
        if self.min_row == 1 and self.max_row == MAX_ROW:
            return f"{get_column_letter(self.min_col)}:{get_column_letter(self.max_col)}"
        if self.min_col == 1 and self.max_col == MAX_COL:
            return f"{self.min_row}:{self.max_row}"
        start = a1(self.min_row, self.min_col)
        return start if self.size == 1 else f"{start}:{a1(self.max_row, self.max_col)}"


def a1(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


@lru_cache(maxsize=65536)
def parse_area(text: str) -> Area | None:
    """'B3', '$B$3:F10', 'A:C' or '3:5' as an Area, or None when the text is not an A1 reference."""
    if match := _AREA.fullmatch(text):
        cols = sorted((_column(match[1]), _column(match[3] or match[1])))
        rows = sorted((int(match[2]), int(match[4] or match[2])))
    elif match := _COLUMNS.fullmatch(text):
        cols, rows = sorted((_column(match[1]), _column(match[2]))), [1, MAX_ROW]
    elif match := _ROWS.fullmatch(text):
        cols, rows = [1, MAX_COL], sorted((int(match[1]), int(match[2])))
    else:
        return None
    if rows[0] < 1 or rows[1] > MAX_ROW or cols[1] > MAX_COL:
        return None
    return Area(rows[0], cols[0], rows[1], cols[1])


def _column(letters: str) -> int:
    return column_index_from_string(letters.upper())


def quote(sheet: str) -> str:
    """A sheet name as a reference needs it: quoted when it is not a plain word."""
    if _PLAIN.fullmatch(sheet) and not _CELL_LIKE.fullmatch(sheet):
        return sheet
    return "'" + sheet.replace("'", "''") + "'"


def split_ref(text: str) -> tuple[str | None, str | None, str]:
    """Split "'[Book 1.xlsx]Sheet 1'!A1" into its workbook, sheet and reference parts."""
    if text.startswith("'"):
        end = 1
        while (end := text.find("'", end)) != -1 and text[end + 1 : end + 2] == "'":
            end += 2
        if end == -1 or text[end + 1 : end + 2] != "!":
            return None, None, text
        prefix, ref = text[1:end].replace("''", "'"), text[end + 2 :]
    else:
        prefix, bang, ref = text.rpartition("!")
        if not bang:
            return None, None, text
    book = None
    if prefix.startswith("[") and "]" in prefix:
        book, prefix = prefix[1:].rsplit("]", 1)
        if len(prefix) > 1 and prefix[0] == prefix[-1] == "'":  # [book.xlsx]'Sheet 1'!A1, as agents write it
            prefix = prefix[1:-1].replace("''", "'")
    return book, prefix or None, ref


def cite(book: str, sheet: str, ref: str) -> str:
    """A citation in Excel's own form for a reference to another workbook: [model.xlsx]Summary!B4."""
    inner = f"[{book}]{sheet}"
    if _BOOK.fullmatch(book) and quote(sheet) == sheet:
        return f"{inner}!{ref}"
    return "'" + inner.replace("'", "''") + f"'!{ref}"


def formula_refs(formula: str) -> tuple[list[str], str | None]:
    """The references a formula reads, and a warning when that list may be incomplete.

    Formulas filled across a sheet differ only in their cells, so each shape (the formula with every
    cell written A1) is tokenised once, and its cells are put back in order."""
    notes, text = [], _STRING.sub('""', formula)
    if _SPILL.search(text):
        text = _SPILL.sub(r"\1", text)
        notes.append("reads a spilled range (A1#), and only its first cell is known")
    cells = _CELLISH.findall(text)
    shape = _shape(_CELLISH.sub("A1", text), len(cells))
    if shape is None:
        shape = _shape(text, 0)
        cells = []
    if shape is None:
        return [], "could not be parsed"
    templates, functions, deleted = shape
    refs, remaining = [], iter(cells)
    for template, count, keep in templates:
        own = iter([next(remaining) for _ in range(count)])
        if keep:
            refs.append(_CELLISH.sub(lambda _: next(own), template) if count else template)
    if "_XLFN.ANCHORARRAY(" in functions and not notes:
        notes.append("reads a spilled range (A1#), and only its first cell is known")
    if deleted:
        notes.append("reads a deleted range (#REF!)")
    if functions & _DYNAMIC:
        notes.append("builds references with INDIRECT or OFFSET, which are only known when Excel calculates")
    return refs, " and ".join(notes) or None


@lru_cache(maxsize=8192)
def _shape(text, cells):
    """A formula's tokens that hold cells, each with how many and whether it is a reference to keep; the
    functions it calls; and whether it reads a deleted range. None when it cannot be parsed, or when its
    tokens do not hold the number of cells expected."""
    try:
        tokens = Tokenizer(text).items
    except Exception:  # the tokenizer raises IndexError as well as TokenizerError on a malformed formula
        return None
    templates, held = [], 0
    for token in tokens:
        count = len(_CELLISH.findall(token.value)) if cells else 0
        held += count
        if token.type == "OPERAND" or count:
            value = token.value
            keep = token.type == "OPERAND" and token.subtype == "RANGE"
            keep = keep and not value.startswith("_xlpm.") and "#REF!" not in value.upper()
            templates.append((value, count, keep))
    if held != cells:
        return None
    functions = frozenset(t.value.upper() for t in tokens if t.type == "FUNC")
    deleted = any(t.type == "OPERAND" and "#REF!" in t.value.upper() for t in tokens)
    return tuple(templates), functions, deleted


def excel_formula(formula: str) -> str:
    """A formula as Excel shows it, without the prefixes the file format adds: _xlfn.XLOOKUP is XLOOKUP,
    _xlfn.ANCHORARRAY(B2) is B2# and _xlfn.SINGLE(A1:A9) is @A1:A9. Text in quotes is left alone."""
    parts = _QUOTED.split(formula)
    for i in range(0, len(parts), 2):
        parts[i] = _PREFIX.sub("", _SINGLE.sub(r"@\1", _ANCHOR.sub(r"\1#", parts[i])))
    return "".join(parts)
