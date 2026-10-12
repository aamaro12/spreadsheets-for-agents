"""Each reading step on saved workbooks: what it shows, what it cites and how it fails."""

from datetime import datetime
import re
from unittest.mock import patch
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table

from sheetscope import SheetError, load
from sheetscope.render import markdown
from support import WorkbookCase


def walk(node):
    """Every node of a trace tree, parents first."""
    yield node
    for child in node.get("children", []):
        yield from walk(child)


def refs(result):
    return [f"{m['sheet']}!{m['cell']}" for m in result["matches"]]


class LoadTests(WorkbookCase):
    def test_files_that_are_not_workbooks_are_refused_with_a_reason(self):
        old = self.folder / "old.xlsx"
        old.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64))
        broken = self.folder / "broken.xlsx"
        broken.write_bytes(b"PK\x03\x04 not a zip archive")
        other = self.folder / "other.xlsx"
        with zipfile.ZipFile(other, "w") as archive:
            archive.writestr("readme.txt", "not a workbook")
        lock = self.folder / "~$model.xlsx"
        lock.write_bytes(b"\x0bAlessandro" + bytes(154))
        cases = {
            lock: "the lock file Excel keeps while model.xlsx is open",
            self.folder / "missing.xlsx": "No such workbook",
            self.folder / "notes.csv": "not an .xlsx or .xlsm workbook",
            old: "password-protected or in the old .xls format",
            broken: "damaged",
            other: "not a workbook: it has no workbook part",
        }
        for path, reason in cases.items():
            with self.subTest(path.name), self.assertRaisesRegex(SheetError, reason):
                load(path)

    def test_xml_that_declares_entities_is_refused_in_any_part(self):
        for part, root in (
            ("xl/worksheets/sheet1.xml", "worksheet"),
            ("xl/workbook.xml", "workbook"),
            ("xl/styles.xml", "styleSheet"),
        ):
            path = self.save(Workbook(), name=f"{root}.xlsx")
            with zipfile.ZipFile(path) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            bomb = f'<!DOCTYPE {root} [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;">]><{root}'
            parts[part] = parts[part].decode().replace(f"<{root}", bomb, 1).encode()
            with zipfile.ZipFile(path, "w") as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            with self.subTest(part=part), self.assertRaisesRegex(SheetError, "refused for safety"):
                load(path)

    def test_xml_in_an_encoding_that_could_hide_a_declaration_is_refused(self):
        path = self.save(Workbook())
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        styles = parts["xl/styles.xml"].decode()
        for encoding, refused in (("cp037", True), ("ISO-8859-1", False)):
            padding = " " * 300  # a declaration may hold any amount of space before its encoding
            parts["xl/styles.xml"] = f'<?xml version="1.0"{padding}encoding="{encoding}"?>{styles}'.encode()
            with zipfile.ZipFile(path, "w") as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            with self.subTest(encoding=encoding):
                if refused:
                    with self.assertRaisesRegex(SheetError, "cp037 encoding, which is refused for safety"):
                        load(path)
                else:
                    self.assertEqual(load(path).name, "book.xlsx")

    def test_a_workbook_that_unpacks_too_large_is_refused(self):
        path = self.save(Workbook())
        with patch("sheetscope.book.MAX_UNPACKED", 1000), self.assertRaisesRegex(SheetError, "over the .* limit"):
            load(path)


class RobustnessTests(WorkbookCase):
    def test_a_malformed_formula_does_not_stop_the_rest(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Model"
        sheet["B1"], sheet["B2"], sheet["C1"], sheet["B9"] = 1, 2, "=B1+B2", "=SUM(B1:B2))"
        loaded = load(self.save(book, {"Model": {"C1": 3, "B9": 3}}))
        self.assertEqual(loaded.cell("Model!C1")["precedents"][0]["ref"], "Model!B1")
        self.assertIn("could not be parsed", " ".join(loaded.cell("Model!B9")["notes"]))
        self.assertEqual(len(loaded.read("Model")["formulas"]), 2)

    def test_cells_filled_by_an_array_formula_show_it(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Calc"
        for row, (units, price) in enumerate([(1, 10), (2, 20), (3, 30)], 2):
            sheet.cell(row, 2, units)
            sheet.cell(row, 3, price)
        sheet["D2"] = ArrayFormula("D2:D4", "=B2:B4*C2:C4")
        sheet["D3"], sheet["D4"] = 40, 90
        loaded = load(self.save(book, {"Calc": {"D2": 10}}))
        member = loaded.cell("Calc!D3")
        self.assertEqual((member["formula"], member["array"], member["display"]), ("=B2:B4*C2:C4", "D2:D4", "40"))
        self.assertEqual([item["ref"] for item in member["precedents"]], ["Calc!B2:B4", "Calc!C2:C4"])
        self.assertEqual(
            loaded.read("Calc")["formulas"], [{"range": "D2:D4", "formula": "=B2:B4*C2:C4", "array": True}]
        )

    def test_newer_functions_show_as_excel_shows_them(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "S"
        for row in (["North", 10], ["South", 20]):
            sheet.append(row)
        sheet["D1"], sheet["D2"] = "South", "=_xlfn.XLOOKUP(D1,A1:A2,B1:B2)"
        loaded = load(self.save(book, {"S": {"D2": 20}}))
        shown = "=XLOOKUP(D1,A1:A2,B1:B2)"
        self.assertEqual(loaded.cell("S!D2")["formula"], shown)
        self.assertEqual(loaded.read("S!D2")["formulas"], [{"range": "D2", "formula": shown}])
        self.assertEqual(refs(loaded.find("xlookup(d1", formulas=True)), ["S!D2"])
        self.assertEqual(loaded.trace("S!D2")["root"]["formula"], shown)

    def test_an_escaped_column_name_is_read_as_a_column(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        for row in (["Region", "# of units"], ["North", 5], ["South", 7]):
            sheet.append(row)
        sheet.add_table(Table(displayName="Sales", ref="A1:B3"))
        sheet["D1"] = "=SUM(Sales['# of units])"
        loaded = load(self.save(book, {"Data": {"D1": 12}}))
        self.assertEqual(loaded.cell("Data!D1")["precedents"][0]["ref"], "Data!B2:B3")

    def test_a_slice_does_not_take_labels_from_a_table_beside_it(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "T"
        for row in (
            ["Scheme", "Value", None, "Item", "FY2025"],
            ["Pensions", 10, None, "Rent", 5],
            ["Bonds", 20, None, None, 7],
        ):
            sheet.append(row)
        result = load(self.save(book)).read("T!D2:E3")
        self.assertEqual([row.get("label") for row in result["rows"]], [None, None])

    def test_wide_ranges_are_cut_to_the_cells_asked_for(self):
        book = Workbook()
        for row in range(1, 4):
            for col in range(1, 301):
                book.active.cell(row, col, row * col)
        result = load(self.save(book)).read("Sheet", max_cells=60)
        self.assertLessEqual(len(result["cells"]), 60)
        self.assertTrue(result["truncated"])

    def test_a_cut_range_shows_only_rows_and_columns_that_keep_cells(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        sheet["E1"], sheet["A3"] = "Notes", "Title"
        for row in range(5, 40):
            for col in range(2, 5):
                sheet.cell(row, col, row * col)
        loaded = load(self.save(book))
        for size in (2, 4, 6, 9):
            with self.subTest(max_cells=size):
                result = loaded.read("Data", max_cells=size)
                shown = {(cell["cell"][0], int(cell["cell"][1:])) for cell in result["cells"]}
                self.assertTrue(result["truncated"])
                self.assertTrue(0 < len(shown) <= size)
                self.assertEqual({column["col"] for column in result["columns"]}, {col for col, _ in shown})
                self.assertEqual({row["row"] for row in result["rows"]}, {row for _, row in shown})

    def test_formulas_filled_down_and_across_are_one_group(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Grid"
        for row in range(1, 4):
            sheet.cell(row, 1, row)
            for col in (2, 3):
                sheet.cell(row, col, f"={get_column_letter(col - 1)}{row}*2")
        result = load(self.save(book)).read("Grid")
        self.assertEqual(result["formulas"], [{"range": "B1:C3", "formula": "=A1*2"}])

    def test_a_workbook_with_only_a_chart_sheet(self):
        book = Workbook()
        data = book.active
        data.append([1])
        data.append([2])
        chart = BarChart()
        chart.add_data(Reference(data, min_col=1, min_row=1, max_row=2))
        book.create_chartsheet("Chart").add_chart(chart)
        book.remove(data)
        loaded = load(self.save(book))
        self.assertEqual((loaded.outline()["sheets"], loaded.outline()["cite_as"]), ([], ""))
        with self.assertRaisesRegex(SheetError, "has no worksheets"):
            loaded.cell("A1")

    def test_dates_in_a_1904_workbook(self):
        book = Workbook()
        book.epoch = CALENDAR_MAC_1904
        book.active["A1"] = datetime(2025, 3, 31)
        book.active["A1"].number_format = "yyyy-mm-dd"
        self.assertEqual(load(self.save(book)).cell("A1")["display"], "2025-03-31")

    def test_internal_names_are_left_out_and_long_lists_are_cut(self):
        book = Workbook()
        book.active["A1"] = 1
        for i in range(60):
            book.defined_names[f"N{i}"] = DefinedName(f"N{i}", attr_text="Sheet!$A$1")
        book.defined_names["_xlfn.XLOOKUP"] = DefinedName("_xlfn.XLOOKUP", attr_text="#NAME?", hidden=True)
        for i in range(30):
            book.defined_names[f"IQ_{i}"] = DefinedName(f"IQ_{i}", attr_text=str(i), hidden=True)
        outline = load(self.save(book)).outline()
        self.assertEqual((len(outline["names"]), outline["more_names"], outline["hidden_names"]), (50, 10, 30))
        self.assertFalse({"_xlfn.XLOOKUP", "IQ_0"} & {name["name"] for name in outline["names"]})


class OutlineTests(WorkbookCase):
    def setUp(self):
        super().setUp()
        self.outline = load(self.northwind()).outline()

    def test_sheets_with_their_blocks_headers_and_labels(self):
        sheets = {sheet["name"]: sheet for sheet in self.outline["sheets"]}
        self.assertEqual(list(sheets), ["Summary", "Model", "Inputs", "Customers", "Scratch"])
        self.assertEqual(sheets["Scratch"]["state"], "hidden")
        self.assertEqual(sheets["Inputs"]["comments"], 1)
        block = sheets["Model"]["blocks"][0]
        self.assertEqual((block["range"], block["header_row"]), ("A1:E12", 3))
        self.assertEqual(block["title"], "Northwind Trading: operating model")
        self.assertEqual([c["text"] for c in block["columns"]], ["£k", "FY2024", "FY2025", "FY2026", "FY2027"])
        self.assertEqual([label["text"] for label in block["labels"]][:3], ["Revenue", "Growth", "Gross profit"])
        self.assertEqual(self.outline["cite_as"], "[northwind.xlsx]Summary!A1")

    def test_names_with_values_and_tables_with_columns(self):
        growth = next(name for name in self.outline["names"] if name["name"] == "Growth")
        self.assertEqual(growth, {"name": "Growth", "refers_to": "Inputs!$B$4", "value": "10.0%"})
        table = {
            "name": "Customers",
            "sheet": "Customers",
            "range": "A1:C6",
            "columns": ["Customer", "Region", "Revenue FY2024"],
        }
        self.assertEqual(self.outline["tables"], [table])
        self.assertEqual(self.outline["warnings"], [])


class OutlineSizeTests(WorkbookCase):
    def test_an_outline_of_many_sheets_stays_short(self):
        book = Workbook()
        book.remove(book.active)
        for number in range(60):
            sheet = book.create_sheet(f"Region {number}")
            sheet.append(["£k", *[f"FY{year}" for year in range(2020, 2032)]])
            for row in range(30):
                sheet.append([f"Line item {row}", *[row * year for year in range(12)]])
            sheet.append([])
            sheet.append([])
            sheet.append(["Notes", "Source: management accounts"])
        outline = load(self.save(book)).outline()
        self.assertEqual(len(outline["sheets"]), 60)
        self.assertLess(len(markdown(outline)), 20_000)


class WarningTests(WorkbookCase):
    def test_formulas_saved_without_values_are_flagged(self):
        book = Workbook()
        book.active["A1"] = 2
        book.active["B1"] = "=A1*2"
        path = self.folder / "unsaved.xlsx"
        book.save(path)
        loaded = load(path)
        outline = loaded.outline()
        self.assertEqual(outline["sheets"][0]["unsaved"], 1)
        self.assertIn("1 formula has no saved value", outline["warnings"][0])
        cell = loaded.cell("B1")
        self.assertEqual((cell["saved"], cell["display"], cell["formula"]), (False, "", "=A1*2"))

    def test_errors_and_hidden_rows_and_columns_are_reported(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        sheet["A1"], sheet["B1"], sheet["A4"] = 1, "=A1/0", "Hidden row"
        for row in (3, 4, 5):
            sheet.row_dimensions[row].hidden = True
        sheet.column_dimensions["C"].hidden = True
        sheet.column_dimensions.group("H", "J", hidden=True)
        loaded = load(self.save(book, {"Data": {"B1": "#DIV/0!"}}))
        data = loaded.outline()["sheets"][0]
        self.assertEqual((data["errors"], data["hidden_rows"], data["hidden_columns"]), (1, "3:5", "C, H:J"))
        error = loaded.cell("B1")
        self.assertEqual((error["type"], error["display"]), ("error", "#DIV/0!"))
        self.assertEqual(loaded.read("Data")["rows"][-1], {"row": 4, "hidden": True})


class FindTests(WorkbookCase):
    def setUp(self):
        super().setUp()
        self.book = load(self.northwind())

    def test_exact_labels_come_first_and_hidden_sheets_are_flagged(self):
        result = self.book.find("Revenue")
        self.assertEqual(refs(result)[:2], ["Summary!A4", "Model!A4"])
        hidden = next(match for match in result["matches"] if match["sheet"] == "Scratch")
        self.assertEqual(hidden["hidden"], "sheet")

    def test_numbers_match_what_they_round_to(self):
        self.assertEqual(refs(self.book.find("1,597")), ["Summary!B4", "Model!E4"])
        self.assertEqual(refs(self.book.find("1597.2")), ["Summary!B4", "Model!E4"])
        self.assertIn("Inputs!B4", refs(self.book.find("10%")))
        self.assertEqual(refs(self.book.find("(49)")), ["Model!B11"])
        match = self.book.find("1,452")["matches"][0]
        expected = ("Revenue", "FY2026", "[northwind.xlsx]Model!D4")
        self.assertEqual((match["label"], match["header"], match["citation"]), expected)
        self.assertNotIn("Model!E4", refs(self.book.find("59")))  # 1,597 does not round to 59

    def test_numbers_match_what_cells_show_in_thousands_millions_or_percent(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Revenue", 1597000, 1600000, 0.125, -1597000, 1597.2])
        formats = ("#,##0,", '#,##0.0,,"m"', "0.0%", "#,##0,;(#,##0,)", "#,##0")
        for cell, fmt in zip(sheet[1][1:], formats):
            cell.number_format = fmt
        loaded = load(self.save(book))
        cases = {"1,597": ["Sheet!B1", "Sheet!F1"], "1.6": ["Sheet!C1"], "12.5": ["Sheet!D1"], "(1,597)": ["Sheet!E1"]}
        for query, cells in cases.items():
            with self.subTest(query=query):
                self.assertEqual(sorted(refs(loaded.find(query))), cells)
        self.assertEqual(refs(loaded.find("59")), [])

    def test_text_a_number_format_adds_is_found(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["EV / EBITDA", 7.04, 2024, 1500])
        sheet["B1"].number_format, sheet["C1"].number_format, sheet["D1"].number_format = '0.0"x"', '"FY"0', "#,##0"
        loaded = load(self.save(book))
        self.assertEqual(refs(loaded.find("7.0x")), ["Sheet!B1"])
        self.assertEqual(refs(loaded.find("fy2024")), ["Sheet!C1"])
        self.assertEqual(refs(loaded.find("ebitda")), ["Sheet!A1"])

    def test_formula_text_can_be_searched(self):
        self.assertEqual(refs(self.book.find("growth", formulas=True)), ["Model!C4", "Model!D4", "Model!E4"])

    def test_limits_and_sheet_filters(self):
        result = self.book.find("e", limit=2)
        self.assertEqual(len(result["matches"]), 2)
        self.assertGreater(result["total"], 2)
        self.assertEqual({match["sheet"] for match in self.book.find("e", sheet="inputs")["matches"]}, {"Inputs"})
        with self.assertRaisesRegex(SheetError, "Did you mean Model"):
            self.book.find("e", sheet="Modle")
        with self.assertRaisesRegex(SheetError, "Give some text"):
            self.book.find("  ")


class ReadTests(WorkbookCase):
    def setUp(self):
        super().setUp()
        self.book = load(self.northwind())

    def test_a_slice_keeps_its_headers_and_labels(self):
        result = self.book.read("Model!C4:E8")
        self.assertEqual([column.get("header") for column in result["columns"]], ["FY2025", "FY2026", "FY2027"])
        labels = ["Revenue", "Growth", "Gross profit", "Operating costs", "EBIT"]
        self.assertEqual([row.get("label") for row in result["rows"]], labels)
        self.assertEqual(result["citation"], "[northwind.xlsx]Model!C4:E8")

    def test_a_sheet_reads_whole_with_its_formulas_grouped(self):
        result = self.book.read("model")
        self.assertEqual(result["range"], "A1:E12")
        self.assertEqual([row["row"] for row in result["rows"]], [1, 3, 4, 5, 6, 7, 8, 9, 11, 12])
        self.assertIn({"range": "C4:E4", "formula": "=B4*(1+Growth)"}, result["formulas"])
        self.assertIn({"range": "B8:E8", "formula": "=B6-B7"}, result["formulas"])
        self.assertEqual(result["merged"], ["A1:E1"])
        shown = {cell["cell"]: cell["display"] for cell in result["cells"]}
        self.assertEqual((shown["B11"], shown["C5"], shown["E4"]), ("(49)", "10.0%", "1,597"))

    def test_names_tables_and_citations_are_references_too(self):
        growth = self.book.read("Growth")
        self.assertEqual((growth["sheet"], growth["range"]), ("Inputs", "B4"))
        self.assertEqual(self.book.read("Customers[]")["range"], "A2:C6")
        self.assertEqual(self.book.read("Customers[#All]")["range"], "A1:C6")
        self.assertEqual(self.book.read("Customers[Revenue FY2024]")["range"], "C2:C6")
        self.assertEqual(self.book.read("[northwind.xlsx]Model!B4")["range"], "B4")

    def test_unclear_or_unknown_references_get_a_useful_error(self):
        cases = {
            "B4": "Name the sheet, for example Summary!B4",
            "Growht": "Did you mean Growth?",
            "Modle!A1": "there is no sheet called Modle. Did you mean Model?",
            "[other.xlsx]Model!B4": "cites other.xlsx",
            "Customers[Profit]": "Customers has no column called Profit",
            "": "Give a reference",
        }
        for ref, message in cases.items():
            with self.subTest(ref=ref), self.assertRaisesRegex(SheetError, re.escape(message)):
                self.book.read(ref)

    def test_large_ranges_are_cut_to_whole_rows(self):
        result = self.book.read("Model", max_cells=12)
        self.assertTrue(result["truncated"])
        self.assertEqual([row["row"] for row in result["rows"]], [1, 3])


class CellTests(WorkbookCase):
    def setUp(self):
        super().setUp()
        self.book = load(self.northwind())

    def test_a_formula_cell_shows_what_it_reads_and_what_uses_it(self):
        cell = self.book.cell("Model!D4")
        self.assertEqual(
            (cell["display"], cell["value"], cell["formula"]), ("1,452", 1452.0000000000002, "=C4*(1+Growth)")
        )
        self.assertEqual((cell["label"], cell["header"], cell["number_format"]), ("Revenue", "FY2026", "#,##0;(#,##0)"))
        reads = [(item["ref"], item.get("via"), item["display"]) for item in cell["precedents"]]
        self.assertEqual(reads, [("Model!C4", None, "1,320"), ("Inputs!B4", "Growth", "10.0%")])
        self.assertEqual(cell["dependents"], ["Summary!D4", "Model!E4", "Model!D5", "Model!E5", "Model!D6", "Model!D9"])

    def test_an_input_shows_its_comment_and_users(self):
        cell = self.book.cell("Growth")
        self.assertEqual((cell["sheet"], cell["cell"], cell["comment"]), ("Inputs", "B4", "Board plan, March 2025"))
        self.assertEqual(cell["dependents"], ["Model!C4", "Model!D4", "Model!E4"])
        self.assertNotIn("precedents", cell)

    def test_one_cell_at_a_time(self):
        with self.assertRaisesRegex(SheetError, "give one cell"):
            self.book.cell("Model!B4:C4")
        merged = self.book.cell("Model!C1")
        self.assertEqual((merged["merged"], merged["type"]), ("A1:E1", "empty"))


class ReferenceTests(WorkbookCase):
    """Formulas that read other sheets, tables, names, whole columns and other workbooks."""

    def setUp(self):
        super().setUp()
        book = Workbook()
        jan = book.active
        jan.title = "Jan"
        feb, mar, data, pl = (book.create_sheet(name) for name in ("Feb", "Mar", "Data", "P&L 2025"))
        for sheet, amount in ((jan, 10), (feb, 20), (mar, 30)):
            sheet["A2"], sheet["B2"] = "Sales", amount
        jan["C2"] = "=Rate*2"
        jan.defined_names["Rate"] = DefinedName("Rate", attr_text="Jan!$B$2")
        book.defined_names["Vat"] = DefinedName("Vat", attr_text="0.2")
        book.defined_names["Old"] = DefinedName("Old", attr_text="#REF!")
        book.defined_names["Gone"] = DefinedName("Gone", attr_text="Jan!#REF!")
        data.append(["Region", "Amount", "Double"])
        data.append(["North", 5, "=Sales[[#This Row],[Amount]]*2"])
        data.append(["South", 7, "=Sales[[#This Row],[Amount]]*2"])
        data.add_table(Table(displayName="Sales", ref="A1:C3"))
        formulas = {
            "B1": ("=SUM(Jan:Mar!B2)", 60),
            "B2": ("=SUM(Sales[Amount])", 12),
            "B3": ("='P&L 2025'!B1*Vat", 12.0),
            "B4": ("=[1]Prices!A1*2", 0),
            "B5": ('=INDIRECT("B1")', 60),
            "B6": ("=SUM(Jan!B2#)", 10),
            "B7": ("=Jan!Rate", 10),
            "D1": ("=SUM(Jan!B:B)", 10),
            "E1": ("=E2+1", 0),
            "E2": ("=E1+1", 0),
        }
        for ref, (formula, _) in formulas.items():
            pl[ref] = formula
        pl.row_dimensions[4].hidden = True
        pl.column_dimensions["D"].hidden = True
        results = {
            "P&L 2025": {ref: value for ref, (_, value) in formulas.items()},
            "Jan": {"C2": 20},
            "Data": {"C2": 10, "C3": 14},
        }
        self.book = load(self.save(book, results))

    def test_three_dimensional_references_cover_each_sheet(self):
        cell = self.book.cell("'P&L 2025'!B1")
        self.assertEqual([item["ref"] for item in cell["precedents"]], ["Jan!B2", "Feb!B2", "Mar!B2"])
        self.assertIn("'P&L 2025'!B1", self.book.cell("Feb!B2")["dependents"])

    def test_tables_resolve_to_their_columns_and_rows(self):
        self.assertEqual(self.book.cell("Data!C3")["precedents"][0]["ref"], "Data!B3")
        column = self.book.cell("'P&L 2025'!B2")["precedents"][0]
        self.assertEqual((column["ref"], column["via"], column["values"]), ("Data!B2:B3", "Sales[Amount]", ["5", "7"]))
        self.assertEqual(self.book.read("Sales")["range"], "A2:C3")

    def test_names_constants_and_quoted_sheet_names(self):
        cell = self.book.cell("'P&L 2025'!B3")
        self.assertEqual([item["ref"] for item in cell["precedents"]], ["'P&L 2025'!B1"])
        self.assertIn("Vat is the constant =0.2", cell["notes"])
        self.assertEqual(self.book.cell("Jan!C2")["precedents"][0]["ref"], "Jan!B2")
        self.assertEqual(self.book.cell("'P&L 2025'!B7")["precedents"][0]["via"], "Rate")
        self.assertEqual(self.book.read("Jan!Rate")["range"], "B2")
        self.assertEqual(self.book.cell("'[book.xlsx]P&L 2025'!B1")["citation"], "'[book.xlsx]P&L 2025'!B1")

    def test_names_that_are_constants_or_deleted_say_so(self):
        cases = {
            "Vat": "Vat is the constant =0.2",
            "Old": "Old refers to a deleted range",
            "Gone": "Gone refers to a deleted",
        }
        for ref, message in cases.items():
            with self.subTest(ref=ref), self.assertRaises(SheetError) as caught:
                self.book.read(ref)
            self.assertIn(message, str(caught.exception))
            self.assertNotIn("Did you mean", str(caught.exception))

    def test_references_that_cannot_be_followed_say_why(self):
        def notes(ref):
            return " ".join(self.book.cell(ref).get("notes", []))

        self.assertIn("another workbook", notes("'P&L 2025'!B4"))
        self.assertIn("INDIRECT", notes("'P&L 2025'!B5"))
        self.assertIn("spilled range", notes("'P&L 2025'!B6"))

    def test_whole_columns_hidden_cells_and_cycles(self):
        column = self.book.cell("'P&L 2025'!D1")["precedents"][0]
        self.assertEqual((column["ref"], column["cells"]), ("Jan!B:B", 1))
        self.assertIn("'P&L 2025'!D1", self.book.cell("Jan!B2")["dependents"])
        self.assertEqual(self.book.cell("'P&L 2025'!B4")["hidden"], "row")
        self.assertEqual(self.book.cell("'P&L 2025'!D1")["hidden"], "column")
        nodes = list(walk(self.book.trace("'P&L 2025'!E1")["root"]))
        self.assertTrue(any(node.get("circular") for node in nodes))


class TraceTests(WorkbookCase):
    def setUp(self):
        super().setUp()
        self.book = load(self.northwind())

    def test_a_trace_reaches_the_inputs(self):
        nodes = list(walk(self.book.trace("Summary!B4")["root"]))
        names = [node["ref"] for node in nodes]
        self.assertEqual(names[:5], ["Summary!B4", "Model!E4", "Model!D4", "Model!C4", "Model!B4"])
        table = next(node for node in nodes if node["ref"] == "Customers!C2:C6")
        self.assertEqual(
            (table["via"], table["values"]), ("Customers[Revenue FY2024]", ["310", "285", "240", "205", "160"])
        )
        self.assertTrue(all(node["via"] == "Growth" for node in nodes if node["ref"] == "Inputs!B4"))

    def test_small_ranges_are_followed_into_their_formulas(self):
        group = self.book.trace("Summary!D6", depth=2)["root"]["children"][0]
        self.assertEqual(group["ref"], "Model!B12:E12")
        self.assertEqual(
            [child["ref"] for child in group["children"]], ["Model!B12", "Model!C12", "Model!D12", "Model!E12"]
        )

    def test_repeats_are_marked_and_the_tree_stays_small(self):
        nodes = list(walk(self.book.trace("Summary!C4")["root"]))
        self.assertTrue(any(node["ref"] == "Model!B4" and node.get("repeat") for node in nodes))
        small = list(walk(self.book.trace("Summary!B4", max_nodes=3)["root"]))
        self.assertEqual(len(small), 3)
        self.assertTrue(small[-1].get("more"))

    def test_large_ranges_are_followed_into_their_first_formulas(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "S"
        for row in range(3, 21):
            sheet.cell(row, 1, row)
            sheet.cell(row, 2, f"=A{row}*2")
        sheet["C1"] = "=SUM(B3:B20)"
        group = load(self.save(book)).trace("S!C1")["root"]["children"][0]
        self.assertEqual((group["ref"], len(group["children"]), group["skipped"]), ("S!B3:B20", 12, 6))

    def test_a_cell_nothing_uses_is_not_cut_short(self):
        children = self.book.trace("Model!E12", dependents=True, depth=1)["root"]["children"]
        self.assertEqual([child["ref"] for child in children], ["Summary!B6", "Summary!C6", "Summary!D6"])
        self.assertFalse(any(child.get("more") for child in children))

    def test_dependents_stop_at_the_depth_asked_for(self):
        children = self.book.trace("Growth", dependents=True, depth=1)["root"]["children"]
        self.assertEqual([child["ref"] for child in children], ["Model!C4", "Model!D4", "Model!E4"])
        self.assertTrue(all(child.get("more") for child in children))


class ValueTests(WorkbookCase):
    def test_dates_and_booleans_keep_their_shown_form(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Dates"
        sheet.append(["Period end", datetime(2025, 3, 31), datetime(2025, 6, 30)])
        sheet.append(["Revenue", 100, 120])
        sheet.append(["Audited", True, False])
        sheet["B1"].number_format = sheet["C1"].number_format = "mmm-yy"
        loaded = load(self.save(book))
        cell = loaded.cell("Dates!C2")
        self.assertEqual((cell["header"], cell["label"]), ("Jun-25", "Revenue"))
        date = loaded.cell("Dates!B1")
        self.assertEqual((date["display"], date["value"], date["type"]), ("Mar-25", "2025-03-31T00:00:00", "date"))
        self.assertEqual(loaded.cell("Dates!B3")["display"], "TRUE")

    def test_cells_holding_only_spaces_read_as_blank(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "BS"
        for row in ([None, "Dec. 31, 2023", "Dec. 31, 2022"], ["Cash", 100, 90], ["Commitments", "\xa0", " "]):
            sheet.append(row)
        sheet.append(["Debt", 50, 40])
        loaded = load(self.save(book))
        self.assertEqual(loaded.cell("BS!C4")["header"], "Dec. 31, 2022")
        self.assertEqual(loaded.cell("BS!B3")["type"], "empty")
        self.assertNotIn("B3", [cell["cell"] for cell in loaded.read("BS")["cells"]])
