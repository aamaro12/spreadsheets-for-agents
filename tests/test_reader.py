"""The one-pass reader on packages written as Excel writes them, including what openpyxl cannot write."""

from datetime import datetime
import time
from unittest.mock import patch
import zipfile

from openpyxl.utils import get_column_letter

from sheetscope import SheetError, load
from sheetscope.render import markdown
from support import WorkbookCase

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE = "http://schemas.openxmlformats.org/package/2006/relationships"
STRICT_MAIN = "http://purl.oclc.org/ooxml/spreadsheetml/main"
STRICT_REL = "http://purl.oclc.org/ooxml/officeDocument/relationships"


class ReaderCase(WorkbookCase):
    def package(self, sheets, strings=(), styles="", book="", names="", parts=None, main=MAIN, rel=REL):
        """Save a workbook whose sheets are given as the XML inside <worksheet>, keyed by sheet name."""
        files = {
            "_rels/.rels": f'<Relationships xmlns="{PACKAGE}"><Relationship Id="r1" Type="{rel}/officeDocument" '
            'Target="xl/workbook.xml"/></Relationships>',
        }
        entries, links = [], []
        for number, (name, body) in enumerate(sheets.items(), 1):
            entries.append(f'<sheet name="{name}" sheetId="{number}" r:id="s{number}"/>')
            links.append(f'<Relationship Id="s{number}" Type="{rel}/worksheet" Target="worksheets/sheet{number}.xml"/>')
            files[f"xl/worksheets/sheet{number}.xml"] = f'<worksheet xmlns="{main}" xmlns:r="{rel}">{body}</worksheet>'
        if strings:
            items = "".join(f"<si>{item}</si>" for item in strings)
            files["xl/sharedStrings.xml"] = f'<sst xmlns="{main}">{items}</sst>'
            links.append(f'<Relationship Id="t1" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/>')
        if styles:
            files["xl/styles.xml"] = f'<styleSheet xmlns="{main}">{styles}</styleSheet>'
            links.append(f'<Relationship Id="t2" Type="{rel}/styles" Target="styles.xml"/>')
        files["xl/workbook.xml"] = (
            f'<workbook xmlns="{main}" xmlns:r="{rel}">{book}<sheets>{"".join(entries)}</sheets>{names}</workbook>'
        )
        files["xl/_rels/workbook.xml.rels"] = f'<Relationships xmlns="{PACKAGE}">{"".join(links)}</Relationships>'
        files.update(parts or {})
        path = self.folder / "raw.xlsx"
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in files.items():
                archive.writestr(name, data)
        return path


class CellTests(ReaderCase):
    def test_shared_formulas_are_written_out_for_each_cell(self):
        rows = "".join(
            f'<row r="{r}"><c r="A{r}"><v>{r}</v></c><c r="B{r}">{formula}<v>{2 * r}</v></c></row>'
            for r, formula in (
                (1, '<f t="shared" ref="B1:B3" si="0">A1*2</f>'),
                (2, '<f t="shared" si="0"/>'),
                (3, '<f t="shared" si="0"/>'),
            )
        )
        book = load(self.package({"S": f"<sheetData>{rows}</sheetData>"}))
        self.assertEqual((book.cell("S!B3")["formula"], book.cell("S!B3")["value"]), ("=A3*2", 6))
        self.assertEqual(book.read("S!B1:B3")["formulas"], [{"range": "B1:B3", "formula": "=A1*2"}])

    def test_a_data_table_shows_its_table_formula_in_every_cell(self):
        rows = (
            '<row r="2"><c r="B2"><v>0.08</v></c></row><row r="3"><c r="B3"><v>0.02</v></c></row>'
            '<row r="5"><c r="C5"><f t="dataTable" ref="C5:D6" dt2D="1" dtr="1" r1="B2" r2="B3"/><v>10</v></c>'
            '<c r="D5"><v>11</v></c></row><row r="6"><c r="C6"><v>12</v></c><c r="D6"><v>13</v></c></row>'
        )
        book = load(self.package({"S": f"<sheetData>{rows}</sheetData>"}))
        cell = book.cell("S!D6")
        self.assertEqual((cell["formula"], cell["array"], cell["value"]), ("=TABLE(B2,B3)", "C5:D6", 13))
        self.assertEqual([item["ref"] for item in cell["precedents"]], ["S!B2", "S!B3"])

    def test_cells_and_rows_may_leave_out_their_references(self):
        rows = '<row><c><v>1</v></c><c><v>2</v></c></row><row><c t="inlineStr"><is><t>x</t></is></c></row>'
        book = load(self.package({"S": f"<sheetData>{rows}</sheetData>"}))
        shown = {cell["cell"]: cell["display"] for cell in book.read("S")["cells"]}
        self.assertEqual(shown, {"A1": "1", "B1": "2", "A2": "x"})

    def test_text_keeps_its_runs_and_escapes_but_not_phonetic_guides(self):
        strings = [
            '<r><t>Net </t></r><r><t>sales</t></r><rPh sb="0" eb="1"><t>ネット</t></rPh>',
            '<t xml:space="preserve">Line_x000D_\nbreak</t>',
        ]
        rows = '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>9</v></c></row>'
        book = load(self.package({"S": f"<sheetData>{rows}</sheetData>"}, strings=strings))
        self.assertEqual(book.cell("S!A1")["value"], "Net sales")
        self.assertEqual(book.cell("S!B1")["value"], "Line\r\nbreak")
        self.assertEqual(book.cell("S!C1")["type"], "empty")  # an index past the end of the shared strings

    def test_types_dates_and_the_1904_calendar(self):
        styles = '<numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy-mm-dd"/></numFmts><cellXfs><xf numFmtId="0"/><xf numFmtId="164"/></cellXfs>'
        rows = (
            '<row r="1"><c r="A1" t="d"><v>2025-03-31T00:00:00</v></c><c r="B1" s="1"><v>44650</v></c>'
            '<c r="C1" t="b"><v>1</v></c><c r="D1" t="e"><v>#N/A</v></c><c r="E1" s="9"><v>3</v></c></row>'
        )
        book = load(
            self.package({"S": f"<sheetData>{rows}</sheetData>"}, styles=styles, book='<workbookPr date1904="1"/>')
        )
        self.assertEqual(book.cell("S!A1")["value"], datetime(2025, 3, 31).isoformat())
        self.assertEqual(book.cell("S!B1")["display"], "2026-03-31")  # 44650 days after 1 January 1904
        self.assertEqual([book.cell(f"S!{c}1")["display"] for c in "CDE"], ["TRUE", "#N/A", "3"])

    def test_merged_ranges_keep_only_their_first_cell(self):
        rows = '<row r="1"><c r="A1" t="inlineStr"><is><t>Title</t></is></c><c r="B1"><v>5</v></c></row>'
        book = load(
            self.package({"S": f'<sheetData>{rows}</sheetData><mergeCells><mergeCell ref="A1:B1"/></mergeCells>'})
        )
        self.assertEqual([cell["cell"] for cell in book.read("S")["cells"]], ["A1"])


class PackageTests(ReaderCase):
    def test_strict_workbooks_are_read(self):
        rows = '<row r="1"><c r="A1"><v>42</v></c></row>'
        book = load(self.package({"S": f"<sheetData>{rows}</sheetData>"}, main=STRICT_MAIN, rel=STRICT_REL))
        self.assertEqual(book.cell("S!A1")["value"], 42)

    def test_a_sheet_may_open_with_markup_from_other_namespaces(self):
        compatible = "http://schemas.openxmlformats.org/markup-compatibility/2006"
        body = (
            f'<mc:AlternateContent xmlns:mc="{compatible}"><mc:Choice Requires="x14"/></mc:AlternateContent>'
            '<sheetData><row r="1"><c r="A1"><v>5</v></c></row></sheetData>'
        )
        self.assertEqual(load(self.package({"S": body})).cell("S!A1")["value"], 5)

    def test_parts_found_whatever_their_case_or_path(self):
        rows = '<row r="1"><c r="A1"><v>7</v></c></row>'
        path = self.package({"S": f"<sheetData>{rows}</sheetData>"})
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        files["XL/Worksheets/Sheet1.XML"] = files.pop("xl/worksheets/sheet1.xml")
        files["xl/_rels/workbook.xml.rels"] = files["xl/_rels/workbook.xml.rels"].replace(
            b'Target="worksheets/sheet1.xml"', b'Target="/xl/worksheets/sheet1.xml"'
        )
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in files.items():
                archive.writestr(name, data)
        self.assertEqual(load(path).cell("S!A1")["value"], 7)

    def test_formulas_name_the_other_workbooks_they_read(self):
        rows = '<row r="1"><c r="A1"><f>[1]Prices!A1*2</f><v>8</v></c></row>'
        link = "file:///C:/Models/Prices%202024.xlsx"
        parts = {
            "xl/externalLinks/externalLink1.xml": f'<externalLink xmlns="{MAIN}"/>',
            "xl/externalLinks/_rels/externalLink1.xml.rels": f'<Relationships xmlns="{PACKAGE}"><Relationship Id="e1" '
            f'Type="{REL}/externalLinkPath" Target="{link}" TargetMode="External"/></Relationships>',
        }
        path = self.package(
            {"S": f"<sheetData>{rows}</sheetData>"},
            names='<externalReferences><externalReference r:id="x1"/></externalReferences>',
            parts=parts,
        )
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        files["xl/_rels/workbook.xml.rels"] = files["xl/_rels/workbook.xml.rels"].replace(
            b"</Relationships>",
            f'<Relationship Id="x1" Type="{REL}/externalLink" Target="externalLinks/externalLink1.xml"/></Relationships>'.encode(),
        )
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in files.items():
                archive.writestr(name, data)
        self.assertIn("another workbook (Prices 2024.xlsx)", " ".join(load(path).cell("S!A1")["notes"]))

    def test_comments_links_hidden_columns_and_scoped_names(self):
        rows = '<row r="1"><c r="A1"><v>1</v></c><c r="B1"><v>2</v></c></row>'
        body = (
            f'<cols><col min="3" max="16384" hidden="1"/></cols><sheetData>{rows}</sheetData>'
            '<hyperlinks><hyperlink ref="A1:B1" r:id="h1"/><hyperlink ref="C9" location="S!A1"/></hyperlinks>'
        )
        threaded = (
            "[Threaded comment]\n\nYour version of Excel allows you to read this threaded comment; however, any "
            "edits to it will get removed if the file is opened in a newer version of Excel. Learn more: "
            "https://go.microsoft.com/fwlink/?linkid=870924\n\nComment:\n    Check this figure"
        )
        parts = {
            "xl/worksheets/_rels/sheet1.xml.rels": f'<Relationships xmlns="{PACKAGE}">'
            f'<Relationship Id="h1" Type="{REL}/hyperlink" Target="https://example.com/q3" TargetMode="External"/>'
            f'<Relationship Id="c1" Type="{REL}/comments" Target="../comments1.xml"/></Relationships>',
            "xl/comments1.xml": f'<comments xmlns="{MAIN}"><commentList><comment ref="B1"><text><t>{threaded}</t>'
            "</text></comment></commentList></comments>",
        }
        names = '<definedNames><definedName name="Rate" localSheetId="1">T!$A$1</definedName></definedNames>'
        sheets = {"S": body, "T": '<sheetData><row r="1"><c r="A1"><v>0.2</v></c></row></sheetData>'}
        book = load(self.package(sheets, names=names, parts=parts))
        cell = book.cell("S!B1")
        self.assertEqual((cell["comment"], cell["link"]), ("Check this figure", "https://example.com/q3"))
        self.assertEqual(book.cell("S!C9")["link"], "S!A1")
        self.assertEqual(book.outline()["sheets"][0]["hidden_columns"], "C:XFD")
        self.assertEqual(book.read("T!Rate")["range"], "A1")

    def test_damage_no_spreadsheet_could_save_is_refused_or_contained(self):
        outside = self.package({"S": '<sheetData><row r="1"><c r="ZZZZ1"><v>1</v></c></row></sheetData>'})
        with self.assertRaisesRegex(SheetError, "damaged: a cell has the reference ZZZZ1"):
            load(outside)
        rows = '<row r="1"><c r="A1"><v>1e999</v></c><c r="B1"><v>nan</v></c><c r="C1"><v>2</v></c></row>'
        body = f'<cols><col min="0" max="99999" hidden="1"/></cols><sheetData>{rows}</sheetData>'
        book = load(self.package({"S": body}))
        self.assertEqual(
            [(book.cell(f"S!{c}1")["display"], book.cell(f"S!{c}1")["type"]) for c in "AB"], [("#NUM!", "error")] * 2
        )
        self.assertEqual(book.outline()["sheets"][0]["hidden_columns"], "A:XFD")

    def test_every_part_read_is_checked_whatever_its_name(self):
        bomb = '<!DOCTYPE worksheet [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;">]>'
        sheet = f'{bomb}<worksheet xmlns="{MAIN}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&b;</t></is></c></row></sheetData></worksheet>'
        path = self.package({"S": '<sheetData><row r="1"><c r="A1"><v>1</v></c></row></sheetData>'})
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        unread = {**files, "docProps/app.xml": f"{bomb}<Properties/>".encode()}  # never read, so it does not matter
        renamed = {name: data for name, data in files.items() if name != "xl/worksheets/sheet1.xml"}
        renamed["xl/worksheets/sheet1.bin"] = sheet.encode()
        renamed["xl/_rels/workbook.xml.rels"] = files["xl/_rels/workbook.xml.rels"].replace(
            b"sheet1.xml", b"sheet1.bin"
        )
        for parts, refused in ((unread, False), (renamed, True)):
            with zipfile.ZipFile(path, "w") as archive:
                for name, data in parts.items():
                    archive.writestr(name, data)
            with self.subTest(refused=refused):
                if refused:
                    with self.assertRaisesRegex(SheetError, "document type declaration, which is refused for safety"):
                        load(path)
                else:
                    self.assertEqual(load(path).cell("S!A1")["value"], 1)

    def test_text_a_terminal_cannot_print_stays_escaped(self):
        rows = '<row r="1"><c r="A1" t="s"><v>0</v></c></row>'
        book = load(self.package({"S": f"<sheetData>{rows}</sheetData>"}, strings=["<t>a_xD800_b</t>"]))
        self.assertEqual(book.cell("S!A1")["value"], "a_xD800_b")
        markdown(book.outline()).encode()

    def test_a_name_on_a_chart_sheet_is_left_out(self):
        names = '<definedNames><definedName name="Peak" localSheetId="1">$A$1</definedName></definedNames>'
        path = self.package({"S": "<sheetData/>"}, names=names)
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        files["xl/workbook.xml"] = files["xl/workbook.xml"].replace(
            b"</sheets>", b'<sheet name="Chart1" sheetId="2" r:id="c1"/></sheets>'
        )
        files["xl/_rels/workbook.xml.rels"] = files["xl/_rels/workbook.xml.rels"].replace(
            b"</Relationships>",
            f'<Relationship Id="c1" Type="{REL}/chartsheet" Target="chartsheets/sheet1.xml"/></Relationships>'.encode(),
        )
        files["xl/chartsheets/sheet1.xml"] = f'<chartsheet xmlns="{MAIN}"/>'.encode()
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in files.items():
                archive.writestr(name, data)
        self.assertEqual(load(path).outline()["names"], [])

    def test_damage_built_to_use_up_time_or_memory_is_contained(self):
        with self.assertRaises(SheetError) as refused:
            load(self.package({"S": f'<sheetData><row r="1"><c r="{"A" * 80_000}1"><v>1</v></c></row></sheetData>'}))
        self.assertLess(len(str(refused.exception)), 120)
        merges = "".join(
            f'<mergeCell ref="{a}{r}:{b}{r + 99}"/>'
            for r in range(1, 2000, 100)
            for a, b in (("A", "CV"), ("CW", "GR"))
        )
        corners = '<row r="1"><c r="A1"><v>1</v></c></row><row r="2000"><c r="GR2000"><v>2</v></c></row>'
        links = "".join(f'<hyperlink ref="A{r}:CV{r + 999}" location="S!A1"/>' for r in range(1, 80_000, 1000))
        arrays = "".join(
            f'<row r="{r}"><c r="A{r}"><f t="array" ref="A1:XFD1048576">1</f><v>1</v></c></row>' for r in range(1, 8001)
        )
        repeats = (
            '<row r="1">' + "".join(f'<c r="{get_column_letter(c)}1"><v>{c}</v></c>' for c in range(1, 8001)) + "</row>"
        )
        hidden = '<col min="1" max="16384" hidden="1"/>' * 20_000
        same = '<mergeCell ref="A1:KUL1"/>' * 8000
        sheets = {
            "M": f"<sheetData>{corners}</sheetData><mergeCells>{merges}</mergeCells>",
            "L": f'<sheetData><row r="1"><c r="A1"><v>1</v></c></row></sheetData><hyperlinks>{links}</hyperlinks>',
            "Y": f"<sheetData>{arrays}</sheetData>",
            "R": f"<sheetData>{repeats}</sheetData><mergeCells>{same}</mergeCells>",
            "H": f'<cols>{hidden}</cols><sheetData><row r="1"><c r="A1"><v>1</v></c></row></sheetData>',
        }
        started = time.perf_counter()
        book = load(self.package(sheets))
        book.outline()
        self.assertLess(time.perf_counter() - started, 4)
        self.assertLessEqual(len(book.sheets["L"].links), 101_000)
        self.assertLessEqual(len(book.sheets["M"].layout.anchor), 600_000)
        self.assertEqual(book.cell("Y!A9")["formula"], "=1")

    def test_too_many_cells_and_broken_xml_are_refused_with_a_reason(self):
        rows = '<row r="1"><c r="A1"><v>1</v></c><c r="B1"><v>2</v></c><c r="C1"><v>3</v></c></row>'
        path = self.package({"S": f"<sheetData>{rows}</sheetData>"})
        with patch("sheetscope.book.MAX_CELLS", 2), self.assertRaisesRegex(SheetError, "more than 2 cells"):
            load(path)
        broken = self.package({"S": "<sheetData><row r='1'><c r='A1'><v>1</v></c></sheetData>"})
        with self.assertRaisesRegex(SheetError, "raw.xlsx could not be read: ParseError"):
            load(broken)
