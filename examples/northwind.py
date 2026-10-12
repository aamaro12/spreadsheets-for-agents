"""Build northwind.xlsx, the example workbook: a small operating model for a made-up company.

openpyxl saves formulas without their results, so this script also stores each formula's result,
computed below, the way Excel does when it saves a calculated workbook.

    python examples/northwind.py
"""

import math
from pathlib import Path
import re
import sys
from xml.sax.saxutils import escape
import zipfile

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
_FORMULA_CELL = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)>(<f[^>]*>[^<]*</f>|<f[^>]*/>)(?:<v\s*/>|<v></v>)</c>')


def save_calculated(workbook, path, results):
    """Save a workbook, then store results ({sheet: {cell: value}}) as its formulas' saved values."""
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for number, sheet in enumerate((worksheet.title for worksheet in workbook.worksheets), 1):
        values = results.get(sheet, {})
        part = f"xl/worksheets/sheet{number}.xml"

        def fill(match, values=values):
            ref, attributes, formula = match.groups()
            if ref not in values:
                return match[0]
            result = values[ref]
            attributes = re.sub(r'\s+t="[^"]*"', "", attributes)
            if isinstance(result, bool):
                kind, text = ' t="b"', str(int(result))
            elif isinstance(result, str):
                kind, text = (' t="e"' if result in ERRORS else ' t="str"'), escape(result)
            else:
                kind, text = "", repr(result)
            return f'<c r="{ref}"{attributes}{kind}>{formula}<v>{text}</v></c>'

        parts[part] = _FORMULA_CELL.sub(fill, parts[part].decode()).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return Path(path)


CUSTOMERS = [
    ("Alder Foods", "North", 310),
    ("Birch Retail", "South", 285),
    ("Cedar Hotels", "North", 240),
    ("Damson Cafés", "West", 205),
    ("Elm Wholesale", "South", 160),
]
INPUTS = {"Growth": 0.10, "GrossMargin": 0.42, "Costs": 310, "Inflation": 0.03, "TaxRate": 0.25}
YEARS = ["FY2024", "FY2025", "FY2026", "FY2027"]
COLUMNS = "BCDE"
MONEY, PERCENT = "#,##0;(#,##0)", "0.0%"


def model():
    """The model's figures by year, computed the way its formulas do."""
    revenue = [sum(amount for _, _, amount in CUSTOMERS)]
    costs = [INPUTS["Costs"]]
    for _ in YEARS[1:]:
        revenue.append(revenue[-1] * (1 + INPUTS["Growth"]))
        costs.append(costs[-1] * (1 + INPUTS["Inflation"]))
    gross = [r * INPUTS["GrossMargin"] for r in revenue]
    ebit = [g - c for g, c in zip(gross, costs)]
    tax = [-e * INPUTS["TaxRate"] for e in ebit]
    net = [e + t for e, t in zip(ebit, tax)]
    return {"revenue": revenue, "gross": gross, "costs": costs, "ebit": ebit, "tax": tax, "net": net}


def build(path):
    figures = model()
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    sheet = book.create_sheet("Model")
    inputs = book.create_sheet("Inputs")
    customers = book.create_sheet("Customers")
    scratch = book.create_sheet("Scratch")
    scratch.sheet_state = "hidden"
    results = {"Summary": {}, "Model": {}}

    inputs["A1"] = "Northwind Trading: assumptions"
    inputs.append([])
    inputs.append(["Assumption", "Value", "Unit"])
    rows = [
        ("Revenue growth", "Growth", PERCENT, "a year"),
        ("Gross margin", "GrossMargin", PERCENT, "of revenue"),
        ("Operating costs FY2024", "Costs", MONEY, "£k"),
        ("Cost inflation", "Inflation", PERCENT, "a year"),
        ("Tax rate", "TaxRate", "0%", "of EBIT"),
    ]
    for row, (label, key, fmt, unit) in enumerate(rows, 4):
        inputs.append([label, INPUTS[key], unit])
        inputs.cell(row, 2).number_format = fmt
        if key != "Costs":
            book.defined_names[key] = DefinedName(key, attr_text=f"Inputs!$B${row}")
    inputs["B4"].comment = Comment("Board plan, March 2025", "Finance")

    customers.append(["Customer", "Region", "Revenue FY2024"])
    for row in CUSTOMERS:
        customers.append(list(row))
    customers.add_table(Table(displayName="Customers", ref=f"A1:C{len(CUSTOMERS) + 1}"))

    sheet["A1"] = "Northwind Trading: operating model"
    sheet.merge_cells("A1:E1")
    sheet.append([])
    sheet.append(["£k", *YEARS])
    lines = [
        (4, "Revenue", "revenue", lambda c, p: f"={p}4*(1+Growth)", "=SUM(Customers[Revenue FY2024])"),
        (5, "Growth", None, lambda c, p: f"={c}4/{p}4-1", None),
        (6, "Gross profit", "gross", lambda c, p: f"={c}4*GrossMargin", "=B4*GrossMargin"),
        (7, "Operating costs", "costs", lambda c, p: f"={p}7*(1+Inflation)", "=Inputs!B6"),
        (8, "EBIT", "ebit", lambda c, p: f"={c}6-{c}7", "=B6-B7"),
        (9, "EBIT margin", None, lambda c, p: f"={c}8/{c}4", "=B8/B4"),
        (11, "Tax", "tax", lambda c, p: f"=-{c}8*TaxRate", "=-B8*TaxRate"),
        (12, "Net income", "net", lambda c, p: f"={c}8+{c}11", "=B8+B11"),
    ]
    for row, label, key, formula, first in lines:
        sheet.cell(row, 1, label)
        for i, col in enumerate(COLUMNS):
            text = first if i == 0 else formula(col, COLUMNS[i - 1])
            if text is None:
                continue
            cell = sheet[f"{col}{row}"]
            cell.value = text
            cell.number_format = MONEY if key else PERCENT
            if key:
                results["Model"][cell.coordinate] = figures[key][i]
            elif row == 5:
                results["Model"][cell.coordinate] = figures["revenue"][i] / figures["revenue"][i - 1] - 1
            else:
                results["Model"][cell.coordinate] = figures["ebit"][i] / figures["revenue"][i]

    summary["A1"] = "Northwind Trading: summary"
    summary.append([])
    summary.append(["Metric", "FY2027", "Growth a year from FY2024", "Total FY2024 to FY2027"])
    for row, (label, key, line) in enumerate(
        [("Revenue", "revenue", 4), ("EBIT", "ebit", 8), ("Net income", "net", 12)], 4
    ):
        growth = f"=(Model!E{line}/Model!B{line})^(1/3)-1"
        summary.append([label, f"=Model!E{line}", growth, f"=SUM(Model!B{line}:E{line})"])
        for col, fmt in ((2, MONEY), (3, PERCENT), (4, MONEY)):
            summary.cell(row, col).number_format = fmt
        results["Summary"][f"B{row}"] = figures[key][-1]
        results["Summary"][f"C{row}"] = (figures[key][-1] / figures[key][0]) ** (1 / 3) - 1
        results["Summary"][f"D{row}"] = math.fsum(figures[key])  # the same on every Python version

    scratch["A1"] = "Old scenario, not used"
    scratch.append(["Revenue growth", 0.15])
    return save_calculated(book, path, results)


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_suffix(".xlsx")
    print(build(target))
